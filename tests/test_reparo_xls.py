"""Mecanica de reparo/erros do parser que as fixtures SINTETICAS nao cobrem
(a cobertura principal do reparo/gate/conteudo vive em tests/test_sinteticas.py):

- golden integro NUNCA passa pela rota de reparo;
- arquivo ilegivel -> erro amigavel pedindo re-exportacao (nunca stack cru);
- excecao lendo uma linha -> ErroLeituraExcel com arquivo, linha e conta.

Excecao original do xlrd (antes do
reparo): xlrd.biffh.XLRDError "Expected BOF record; found b'\\x01\\x02\\x06...'"
(b'\\x01\\x02' = BLANK 0x0201 orfao — reproduzida byte a byte na sintetica).
"""
from pathlib import Path

import pytest

from parser.excel import ErroLeituraExcel, ExcelParser, _extrair_stream_biff
from tests.fixtures.gerar_sinteticas import LINHA_PRIMEIRO_LANCAMENTO, adulterar_celula
from tests.helpers import GOLDEN

FIXTURES = Path(__file__).parent / "fixtures"
CORROMPIDA = FIXTURES / "sintetica_corrompida.xls"


def test_golden_integro_nao_passa_pela_rota_de_reparo():
    """Excel integro segue na rota normal do xlrd -- o parser anota que NAO
    houve reparo (a UI so avisa quando houve)."""
    parser = ExcelParser()
    parser.parse(GOLDEN)
    assert parser.reparado is False


def test_reparo_deixa_rastro_no_log_com_o_hash_do_arquivo_e_nao_o_nome(caplog):
    """A UI nao avisa mais o reparo (pro colaborador, reparado com
    self-check verde = arquivo normal); o rastro tecnico e' logging.info no
    parser, com o HASH curto do arquivo (8 hex do sha256 dos bytes), nunca o nome."""
    import hashlib
    import logging
    with caplog.at_level(logging.INFO, logger="parser.excel"):
        ExcelParser().parse(CORROMPIDA)
    mensagens = [r.getMessage() for r in caplog.records]
    h = hashlib.sha256(CORROMPIDA.read_bytes()).hexdigest()[:8]
    assert mensagens == [f"arquivo {h}: reparado na leitura (BOUNDSHEET do exportador do Dominio)"]
    assert "sintetica_corrompida" not in caplog.text


def test_golden_integro_nao_loga_reparo(caplog):
    """Rota normal nao pode poluir o log com falso rastro de reparo."""
    import logging
    with caplog.at_level(logging.INFO, logger="parser.excel"):
        ExcelParser().parse(GOLDEN)
    assert not any("reparad" in r.getMessage() for r in caplog.records)


def test_arquivo_ilegivel_da_erro_amigavel_pedindo_reexportacao(tmp_path):
    """Nem xlrd nem a rota de reparo abrem -> ErroLeituraExcel amigavel com
    arquivo e motivo, nunca stack trace cru na UI."""
    lixo = tmp_path / "nao_e_excel.xls"
    lixo.write_bytes(b"isto nao e um arquivo excel")
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(lixo)
    msg = str(exc.value)
    assert "nao_e_excel.xls" in msg
    assert "não abriu" in msg                       # o fato, sem instrucao


def test_xls_com_ole_marker_danificado_recupera_pela_rota_de_reparo(tmp_path):
    """OLE com o marker little-endian (offset 28) corrompido -> CompDocError no
    xlrd. HOJE vaza CRU (o except so pega XLRDError/IndexError) E perde a
    recuperacao. Com o except largo, a rota de reparo RECUPERA o arquivo inteiro
    -- nunca excecao crua na UI."""
    data = bytearray((FIXTURES / "sintetica_integra.xls").read_bytes())
    data[28:30] = b"\x00\x00"                       # quebra o marker OLE -> CompDocError
    doente = tmp_path / "ole_marker.xls"
    doente.write_bytes(bytes(data))

    parser = ExcelParser()
    contas = parser.parse(doente)                   # hoje: CompDocError cru propaga
    assert parser.reparado is True
    assert len(contas) == 1
    assert contas[0].codigo == "7777"


def test_erro_deliberado_no_reparo_nao_e_re_embrulhado(monkeypatch):
    """Guard-rail do except largo: se um caminho interno do reparo levantar
    ErroLeituraExcel PROPRIA, ela sai LIMPA -- nunca re-embrulhada numa segunda
    camada ('nao abriu: ... nao abriu: ...'). O except ErroLeituraExcel vem
    ANTES do except Exception."""
    marca = "MENSAGEM ORIGINAL LIMPA DO REPARO"

    def _boom(_caminho):
        raise ErroLeituraExcel(marca)

    # CORROMPIDA faz o open estourar (XLRDError) -> entra na rota de reparo,
    # onde _extrair_stream_biff (aqui) levanta a ErroLeituraExcel deliberada.
    monkeypatch.setattr("parser.excel._extrair_stream_biff", _boom)
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(CORROMPIDA)

    msg = str(exc.value)
    assert msg == marca                             # a original, sem embrulho
    assert "nao abriu nem apos" not in msg          # nao foi re-embrulhada
    assert not isinstance(exc.value.__cause__, ErroLeituraExcel)   # sem aninhamento


def test_erro_de_linha_diz_arquivo_linha_e_conta(tmp_path):
    """Qualquer excecao ao ler um lancamento vira ErroLeituraExcel dizendo
    arquivo, linha da planilha e conta em processamento. Fabrica: data do 1o
    lancamento da sintetica corrompida adulterada pra um serial impossivel
    (3e8 -> ano ~823000) -> xldate estoura dentro de _lancamento."""
    stream = _extrair_stream_biff(str(CORROMPIDA))     # ponteiro quebrado preservado
    doente = adulterar_celula(stream, LINHA_PRIMEIRO_LANCAMENTO, 0, 300_000_000)
    arquivo = tmp_path / "sintetica_data_impossivel.xls"
    arquivo.write_bytes(doente)                        # stream BIFF cru vale como .xls
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(arquivo)
    msg = str(exc.value)
    assert "sintetica_data_impossivel.xls" in msg      # arquivo
    assert f"linha {LINHA_PRIMEIRO_LANCAMENTO + 1} da planilha" in msg   # 1-based
    assert "7777" in msg                               # conta em processamento


# === a mensagem de erro diz fato, nao Python ===

def _xls_com_celula_de_debito(caminho, valor, data=None):
    """.xls sintetico no layout do Dominio com UM debito cujo conteudo e' `valor`
    (texto ou float); `data` (serial/float) troca a data desse debito. Dados
    ficticios."""
    from datetime import date

    import xlwt

    from tests.fixtures.gerar_sinteticas import (
        CNPJ, COL_SALDO, EMPRESA, PERIODO, _serial)
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Razão")
    ws.write(0, 0, "Empresa:"); ws.write(0, 2, EMPRESA)
    ws.write(1, 0, "C.N.P.J.:"); ws.write(1, 2, CNPJ)
    ws.write(2, 0, "Período:"); ws.write(2, 2, PERIODO)
    ws.write(3, 0, "RAZÃO")
    ws.write(5, 0, "Data"); ws.write(5, 1, "Número"); ws.write(5, 2, "Histórico")
    ws.write(5, 7, "Cta.C.Part."); ws.write(5, 8, "Débito"); ws.write(5, 9, "Crédito")
    ws.write(5, COL_SALDO, "Saldo-Exercício")
    ws.write(6, 0, "Conta:"); ws.write(6, 1, 7001.0)
    ws.write(6, 2, "2.01.01.001.00"); ws.write(6, 5, "FORNECEDOR FICTICIO LTDA")
    ws.write(7, 2, "SALDO ANTERIOR"); ws.write(7, COL_SALDO, 0.0)
    ws.write(8, 0, _serial(date(2025, 3, 1))); ws.write(8, 1, 5000.0)
    ws.write(8, 2, "Vr. compras conf. 1234 FORNECEDOR FICTICIO LTDA")
    ws.write(8, 7, 38.0); ws.write(8, 9, 100.0); ws.write(8, COL_SALDO, -100.0)
    ws.write(9, 0, _serial(date(2025, 3, 2)) if data is None else data)
    ws.write(9, 1, 5001.0)
    ws.write(9, 2, "PGTO NF 1234"); ws.write(9, 7, 38.0)
    ws.write(9, 8, valor); ws.write(9, COL_SALDO, 0.0)
    ws.write(11, 4, "Total da conta:"); ws.write(11, 8, 100.0); ws.write(11, 9, 100.0)
    wb.save(str(caminho))
    return caminho


@pytest.mark.parametrize("valor,visto", [
    ("abc", "'abc'"),
    ("NaN", "'NaN'"),
    (float("nan"), "'nan'"),
    (1e300, "'1e+300'"),
], ids=["texto", "NaN-texto", "nan-float", "1e300"])
def test_celula_de_dinheiro_ilegivel_vira_fato_com_a_linha(tmp_path, valor, visto):
    """Texto, NaN e valor enorme numa celula de dinheiro -> ErroLeituraExcel
    com arquivo, linha e conta, e a causa em portugues -- nunca a classe Python
    (`[<class 'decimal...'>]`) nem um NaN que passa o parse e estoura no core."""
    arquivo = _xls_com_celula_de_debito(tmp_path / "valor_ruim.xls", valor)
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(arquivo)
    msg = str(exc.value)
    assert msg.startswith("'valor_ruim.xls': erro lendo a linha 10 da planilha, "
                          "conta 7001 (FORNECEDOR FICTICIO LTDA): ")
    assert f"a célula de valor tem {visto}, que não é um valor em reais" in msg
    assert "<class" not in msg


def test_celula_de_dinheiro_com_nan_no_xlsx_tambem_vira_fato(tmp_path):
    """Na rota openpyxl: o TEXTO 'NaN' numa celula de debito do .xlsx."""
    import openpyxl
    wb = openpyxl.load_workbook(FIXTURES / "sintetica_integra.xlsx")
    ws = wb.active
    linha = next(r for r in ws.iter_rows(min_row=1)
                 if isinstance(r[8].value, (int, float)) and r[8].value)
    linha[8].value = "NaN"
    arquivo = tmp_path / "nan.xlsx"
    wb.save(str(arquivo))
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(arquivo)
    msg = str(exc.value)
    assert msg.startswith("'nan.xlsx': erro lendo a linha ")
    assert "a célula de valor tem 'NaN', que não é um valor em reais" in msg
    assert "<class" not in msg


def test_xls_sem_nenhuma_planilha_diz_o_fato(tmp_path):
    """.xls valido sem BOUNDSHEET (id 0x0085 -> 0x0086 no stream) abria
    'Erro inesperado (IndexError)'; agora e' erro de leitura com o nome."""
    stream = bytearray(_extrair_stream_biff(str(FIXTURES / "sintetica_integra.xls")))
    pos = 0
    while pos + 4 <= len(stream):
        rid = stream[pos] | (stream[pos + 1] << 8)
        tam = stream[pos + 2] | (stream[pos + 3] << 8)
        if rid == 0x0085:
            stream[pos] = 0x86
        pos += 4 + tam
    arquivo = tmp_path / "sem_planilha.xls"
    arquivo.write_bytes(bytes(stream))
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(arquivo)
    assert str(exc.value) == "'sem_planilha.xls' não tem nenhuma planilha."


def test_parse_com_nome_cita_o_nome_do_upload_e_sem_nome_o_do_arquivo(tmp_path):
    """`nome=` e' como o arquivo aparece nas mensagens; sem ele, o nome
    do proprio caminho (scripts e testes que chamam parse(caminho))."""
    lixo = tmp_path / "tmpab12cd_lixo.xls"
    lixo.write_bytes(b"isto nao e um arquivo excel")
    with pytest.raises(ErroLeituraExcel) as sem:
        ExcelParser().parse(lixo)
    with pytest.raises(ErroLeituraExcel) as com:
        ExcelParser().parse(lixo, nome="Fornecedor_X.xls")
    assert str(sem.value).startswith("'tmpab12cd_lixo.xls' não abriu nem após")
    assert str(com.value).startswith("'Fornecedor_X.xls' não abriu nem após")


def test_nome_do_upload_vale_em_todas_as_mensagens_do_parser(tmp_path, caplog):
    """Com `nome=`, NENHUMA mensagem volta ao nome do caminho temporario:
    erro de linha, .xlsx quebrado e o rastro do reparo no log."""
    import logging

    ruim = _xls_com_celula_de_debito(tmp_path / "tmpab12cd_valor.xls", "abc")
    with pytest.raises(ErroLeituraExcel) as linha:
        ExcelParser().parse(ruim, nome="Upload.xls")
    assert str(linha.value).startswith("'Upload.xls': erro lendo a linha 10 ")

    zip_quebrado = tmp_path / "tmpab12cd_quebrado.xlsx"
    zip_quebrado.write_bytes(b"PK\x03\x04 isto nao e um zip de verdade")
    with pytest.raises(ErroLeituraExcel) as xlsx:
        ExcelParser().parse(zip_quebrado, nome="Upload.xlsx")
    assert str(xlsx.value).startswith("'Upload.xlsx' é um .xlsx que não abriu")

    with caplog.at_level(logging.INFO, logger="parser.excel"):
        ExcelParser().parse(CORROMPIDA, nome="Upload_reparado.xls")
    reparo = [r.getMessage() for r in caplog.records if "reparad" in r.getMessage()]
    assert reparo and all("Upload_reparado" not in m and "sintetica_corrompida" not in m for m in reparo)   # so o hash


# === o texto cru do leitor (ingles, bytes, Python) vai pro log ===

def _avisos_do_parser(caplog):
    return [r.getMessage() for r in caplog.records
            if r.name == "parser.excel" and r.levelname == "WARNING"]


def test_arquivo_que_nao_abre_nao_mostra_o_texto_do_leitor_e_o_manda_pro_log(
        tmp_path, caplog):
    lixo = tmp_path / "lixo.xls"
    lixo.write_bytes(b"isto nao e um arquivo excel")
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(lixo)
    assert str(exc.value) == "'lixo.xls' não abriu nem após tentativa de reparo."
    avisos = _avisos_do_parser(caplog)
    import hashlib
    assert avisos == [f"arquivo {hashlib.sha256(lixo.read_bytes()).hexdigest()[:8]}: nao abriu (leitor: XLRDError)"]
    assert "lixo.xls" not in avisos[0] and "Unsupported format" not in avisos[0]    # nem nome, nem texto do leitor


def test_linha_ilegivel_sem_causa_conhecida_diz_so_que_nao_foi_reconhecida(
        tmp_path, caplog):
    """Data NaN estoura dentro do leitor ("cannot convert float NaN to integer",
    texto de Python): a tela diz o fato; o texto cru fica no log."""
    arquivo = _xls_com_celula_de_debito(tmp_path / "data_nan.xls", 10.0,
                                        data=float("nan"))
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(arquivo)
    assert str(exc.value) == ("'data_nan.xls': erro lendo a linha 10 da planilha, "
                              "conta 7001 (FORNECEDOR FICTICIO LTDA): "
                              "o conteúdo da linha não foi reconhecido")
    avisos = _avisos_do_parser(caplog)
    import hashlib
    assert avisos == [f"arquivo {hashlib.sha256(arquivo.read_bytes()).hexdigest()[:8]}: linha 10 nao reconhecida "
                      f"(leitor: ValueError)"]
    assert "data_nan" not in avisos[0] and "NaN" not in avisos[0]                  # nem nome, nem texto do Python


def test_celula_de_dinheiro_ilegivel_segue_com_a_frase_propria_e_sem_log(
        tmp_path, caplog):
    """CONTROLE: a causa conhecida (texto em celula de dinheiro) NAO vira o
    generico "nao foi reconhecido" e nao polui o log."""
    arquivo = _xls_com_celula_de_debito(tmp_path / "abc.xls", "abc")
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(arquivo)
    assert str(exc.value).endswith(
        "a célula de valor tem 'abc', que não é um valor em reais")
    assert _avisos_do_parser(caplog) == []
