"""Rota .xlsx do parser: roteamento por MAGIC BYTES + grade openpyxl.

O Dominio tambem exporta .xlsx (ZIP, assinatura PK); as diferencas em relacao ao
.xls sao: datetime nativo na col Data, float/int nas celulas (mesmo perigo de
artefato 1440.090000000001), cabecalho 'Lote' no lugar de 'Numero' e saldo com
sinal negativo. As duas rotas convergem pra MESMA list[Conta]; extensao
mentirosa NUNCA derruba (decide a assinatura, nao o nome). .xlsx corrompido
NAO tem rota de reparo (universo ZIP != defeito BIFF) -> erro amigavel.
"""
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from parser.excel import ErroLeituraExcel, ExcelParser

D = Decimal
FIXTURES = Path(__file__).parent / "fixtures"
INTEGRA_XLS = FIXTURES / "sintetica_integra.xls"


def _escrever_xlsx_dominio(caminho: Path) -> None:
    """Um .xlsx minimo no contrato do Dominio, com os tracos observados no
    arquivo real: header 'Lote', saldo na col 11 com sinal negativo, datetime
    nativo e o artefato de float numa celula de dinheiro."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Razão"
    ws.cell(1, 1, "Empresa:"); ws.cell(1, 3, "EMPRESA FICTICIA LTDA")
    ws.cell(2, 1, "C.N.P.J.:"); ws.cell(2, 3, "00.000.000/0000-00")
    ws.cell(3, 1, "Período:"); ws.cell(3, 3, "01/01/2026 - 31/05/2026")
    ws.cell(5, 1, "RAZÃO")
    ws.cell(7, 1, "Data"); ws.cell(7, 2, "Lote"); ws.cell(7, 3, "Histórico")
    ws.cell(7, 8, "Cta.C.Part."); ws.cell(7, 9, "Débito"); ws.cell(7, 10, "Crédito")
    ws.cell(7, 12, "Saldo-Exercício")
    ws.cell(8, 1, "Conta:"); ws.cell(8, 2, 1201)
    ws.cell(8, 3, "2.01.01.01.10"); ws.cell(8, 6, "FORNECEDOR FICTICIO XLSX LTDA")
    ws.cell(9, 3, "SALDO ANTERIOR"); ws.cell(9, 12, -1000.0)
    ws.cell(10, 1, datetime(2026, 1, 9)); ws.cell(10, 2, 362832)
    ws.cell(10, 3, "Vr. compras conf. 4321 FORNECEDOR FICTICIO XLSX LTDA")
    ws.cell(10, 8, 52); ws.cell(10, 10, 1440.090000000001)
    ws.cell(10, 12, -2440.09)
    ws.cell(11, 1, datetime(2026, 2, 2)); ws.cell(11, 2, 362840)
    ws.cell(11, 3, "Valor ref. a PGTO NF 4321")
    ws.cell(11, 8, 8); ws.cell(11, 9, 500.0)
    ws.cell(11, 12, -1940.09)
    ws.cell(13, 5, "Total da conta:")
    ws.cell(13, 9, 500.0); ws.cell(13, 10, 1440.09)
    wb.save(caminho)


def test_xlsx_feliz_normaliza_data_dinheiro_e_sinal(tmp_path):
    caminho = tmp_path / "razao.xlsx"
    _escrever_xlsx_dominio(caminho)
    parser = ExcelParser()
    contas = parser.parse(caminho)

    assert parser.reparado is False           # reparo e' universo BIFF, nunca xlsx
    assert len(contas) == 1
    conta = contas[0]
    assert conta.codigo == "1201"             # int na celula -> id sem '.0'
    assert conta.nome == "FORNECEDOR FICTICIO XLSX LTDA"
    assert conta.saldo_anterior == D("-1000.00")      # sinal negativo, col movel 11
    assert conta.data_fim_periodo == date(2026, 5, 31)
    assert len(conta.lancamentos) == 2
    compra, pgto = conta.lancamentos
    assert compra.data == date(2026, 1, 9)    # datetime nativo -> date interno
    assert compra.numero == "362832"
    assert compra.credito == D("1440.09")      # artefato de float MORTO
    assert compra.saldo_exercicio == D("-2440.09")
    assert pgto.debito == D("500.00")
    assert conta.total_debito_dominio == D("500.00")
    assert conta.total_credito_dominio == D("1440.09")
    assert conta.saldo_final == D("-1940.09")


def _escrever_xlsx_2contas(caminho: Path) -> None:
    """Duas contas empilhadas no MESMO arquivo, com saldos/totais independentes
    -- o formato da razao empilha N fornecedores e a triagem do lote
    depende disso."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Razão"
    ws.cell(1, 1, "Período:"); ws.cell(1, 3, "01/01/2026 - 31/05/2026")
    ws.cell(3, 1, "Data"); ws.cell(3, 2, "Lote"); ws.cell(3, 3, "Histórico")
    ws.cell(3, 8, "Cta.C.Part."); ws.cell(3, 9, "Débito"); ws.cell(3, 10, "Crédito")
    ws.cell(3, 12, "Saldo-Exercício")
    # Conta A
    ws.cell(5, 1, "Conta:"); ws.cell(5, 2, 1201); ws.cell(5, 3, "2.01.01")
    ws.cell(5, 6, "FORNECEDOR FICTICIO A LTDA")
    ws.cell(6, 3, "SALDO ANTERIOR"); ws.cell(6, 12, -1000.0)
    ws.cell(7, 1, datetime(2026, 1, 10)); ws.cell(7, 2, 1)
    ws.cell(7, 3, "Vr. compras conf. 4321 FORN A"); ws.cell(7, 10, 500.0)
    ws.cell(7, 12, -1500.0)
    ws.cell(8, 5, "Total da conta:"); ws.cell(8, 9, 0.0); ws.cell(8, 10, 500.0)
    # Conta B
    ws.cell(10, 1, "Conta:"); ws.cell(10, 2, 1202); ws.cell(10, 3, "2.01.02")
    ws.cell(10, 6, "FORNECEDOR FICTICIO B LTDA")
    ws.cell(11, 3, "SALDO ANTERIOR"); ws.cell(11, 12, -2000.0)
    ws.cell(12, 1, datetime(2026, 2, 5)); ws.cell(12, 2, 2)
    ws.cell(12, 3, "Vr. compras conf. 8765 FORN B"); ws.cell(12, 10, 300.0)
    ws.cell(12, 12, -2300.0)
    ws.cell(13, 5, "Total da conta:"); ws.cell(13, 9, 0.0); ws.cell(13, 10, 300.0)
    wb.save(caminho)


def test_parser_le_multiplas_contas_no_mesmo_arquivo(tmp_path):
    """parse() devolve UMA Conta por bloco 'Conta:', com saldo anterior, totais
    e lancamentos INDEPENDENTES entre as contas."""
    caminho = tmp_path / "duas_contas.xlsx"
    _escrever_xlsx_2contas(caminho)
    contas = ExcelParser().parse(caminho)

    assert len(contas) == 2
    a, b = contas
    assert a.codigo == "1201" and a.nome == "FORNECEDOR FICTICIO A LTDA"
    assert a.saldo_anterior == D("-1000.00")
    assert a.total_credito_dominio == D("500.00")
    assert a.saldo_final == D("-1500.00")
    assert len(a.lancamentos) == 1
    assert b.codigo == "1202" and b.nome == "FORNECEDOR FICTICIO B LTDA"
    assert b.saldo_anterior == D("-2000.00")     # nao vazou o saldo da conta A
    assert b.total_credito_dominio == D("300.00")
    assert b.saldo_final == D("-2300.00")
    assert len(b.lancamentos) == 1


def test_extensao_mentirosa_bytes_xlsx_com_nome_xls(tmp_path):
    """A rota e' dos MAGIC BYTES (PK), nao da extensao."""
    caminho = tmp_path / "mentira.xls"
    _escrever_xlsx_dominio(caminho)
    contas = ExcelParser().parse(caminho)
    assert contas[0].codigo == "1201"


def test_extensao_mentirosa_bytes_xls_com_nome_xlsx(tmp_path):
    """BIFF nomeado .xlsx cai na rota xlrd normalmente."""
    caminho = tmp_path / "mentira.xlsx"
    caminho.write_bytes(INTEGRA_XLS.read_bytes())
    parser = ExcelParser()
    contas = parser.parse(caminho)
    assert parser.reparado is False
    assert contas[0].codigo == "7777"


def test_xlsx_corrompido_sem_rota_de_reparo_erro_amigavel(tmp_path):
    """ZIP quebrado nao tem reparo (problema diferente do BIFF): erro direto
    pedindo re-exportacao, mensagem pra humano."""
    integro = tmp_path / "ok.xlsx"
    _escrever_xlsx_dominio(integro)
    corrompido = tmp_path / "quebrado.xlsx"
    corrompido.write_bytes(integro.read_bytes()[:200])   # PK presente, ZIP truncado
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(corrompido)
    msg = str(exc.value)
    assert "quebrado.xlsx" in msg
    assert "não abriu" in msg                       # o fato, sem instrucao


def test_xlsx_corrompido_nao_mostra_o_texto_do_leitor_e_o_manda_pro_log(tmp_path, caplog):
    """O texto em ingles do leitor ("File is not a zip file") vai pro log."""
    integro = tmp_path / "ok.xlsx"
    _escrever_xlsx_dominio(integro)
    corrompido = tmp_path / "quebrado.xlsx"
    corrompido.write_bytes(integro.read_bytes()[:200])
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(corrompido)
    assert str(exc.value) == ("'quebrado.xlsx' é um .xlsx que não abriu "
                              "(arquivo corrompido ou incompleto).")
    avisos = [r.getMessage() for r in caplog.records
              if r.name == "parser.excel" and r.levelname == "WARNING"]
    import hashlib
    h = hashlib.sha256(corrompido.read_bytes()).hexdigest()[:8]
    assert avisos == [f"arquivo {h}: nao abriu como .xlsx (leitor: BadZipFile)"]    # hash e tipo; nem nome nem texto do leitor
    assert "quebrado" not in avisos[0] and "zip file" not in avisos[0].lower()


def test_lixo_sem_assinatura_erro_amigavel(tmp_path):
    caminho = tmp_path / "lixo.xlsx"
    caminho.write_bytes(b"\x00\x01\x02isto nao e excel nenhum")
    with pytest.raises(ErroLeituraExcel) as exc:
        ExcelParser().parse(caminho)
    assert "não abriu" in str(exc.value)            # o fato, sem instrucao
