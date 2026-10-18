"""Clareza da tela. Usuario-alvo: colaborador SEM experiencia contabil.
So' a casca (app/) e os textos mudam: os numeros vem prontos do core.
"""
from datetime import date
from decimal import Decimal

import pytest
from streamlit.testing.v1 import AppTest

from app import apresentacao as ap
from core.conferencia import conferir
from core.modelo import Conta
from parser.historico import ClassificadorHistorico
from tests.helpers import GOLDEN, lanc
from tests.test_app_ponte import MAIN
from tests.test_ponte import COMPRA, PAGAMENTO, TRANSFERENCIA, _conta
from tests.test_recebimento import PIX, _pago_a_maior_e_recebimento_sintetico

D = Decimal
CLASSIFICADOR = ClassificadorHistorico()
MISTA = GOLDEN.parent / "sintetica_integra.xls"
PGTO_NF = "- PAGTO REF. NF.:  {nf}.1 - PARC.:FORN"


def _so_credito():
    """Compra de 2.941,92 em aberto e um recebimento de 897,60 no balde."""
    return conferir(Conta(codigo="9000", classificacao="1.1", nome="FORN",
                          data_fim_periodo=date(2025, 12, 31),
                          lancamentos=[
                              lanc("Vr. compras conf. 810112 FORN", numero="c1",
                                   data=date(2025, 12, 11), credito="2941.92",
                                   saldo="-2941.92"),
                              lanc(PIX, numero="810117", data=date(2025, 12, 17),
                                   credito="897.60", saldo="-3839.52")]), CLASSIFICADOR)


def _so_debito():
    """Compra de 1.000,00 em aberto e um pagamento sem NF de 300,00 que nao
    casa com nada (balde so' com debito)."""
    return conferir(_conta([(COMPRA.format(nf="5001"), "0.00", "1000.00", 2),
                            (PAGAMENTO, "300.00", "0.00", 5)]), CLASSIFICADOR)


def _dois_juros():
    """Duas NFs quitadas acima do valor (possivel juros) e uma em aberto."""
    return conferir(_conta([(COMPRA.format(nf="5001"), "0.00", "100.00", 2),
                            (PGTO_NF.format(nf="5001"), "102.46", "0.00", 5),
                            (COMPRA.format(nf="5002"), "0.00", "200.00", 6),
                            (PGTO_NF.format(nf="5002"), "210.00", "0.00", 8),
                            (COMPRA.format(nf="5003"), "0.00", "50.00", 9)]),
                    CLASSIFICADOR)


# --- 1. ponte curta logo abaixo do total -------------------------------------

def test_credito_sem_destino_vira_a_ponte_curta_ate_o_dominio():
    assert ap.ponte_curta_do_total(_so_credito()) == (
        "+ crédito sem destino R$ 897,60 = saldo no Domínio R$ 3.839,52")


def test_debito_sem_destino_entra_com_menos():
    assert ap.ponte_curta_do_total(_so_debito()) == (
        "- débito sem destino R$ 300,00 = saldo no Domínio R$ 700,00")


def test_sem_balde_sem_ponte_curta():
    res = conferir(_conta([(COMPRA.format(nf="5001"), "0.00", "100.00", 2)]),
                   CLASSIFICADOR)
    assert ap.ponte_curta_do_total(res) is None


def test_balde_misto_mostra_os_dois_lados():
    from app.pipeline import processar_excel
    (res,) = processar_excel(MISTA).resultados
    assert ap.ponte_curta_do_total(res) == (
        "- débitos sem destino R$ 3.400,00 + crédito sem destino R$ 500,00 "
        "= saldo no Domínio R$ 5.100,00")


def _tela(caminho):
    at = AppTest.from_file(str(MAIN), default_timeout=60).run()
    at.file_uploader[0].set_value(
        (caminho.name, caminho.read_bytes(), "application/vnd.ms-excel")).run()
    return at


def test_ponte_curta_aparece_na_tela_escapada_contra_lista():
    at = _tela(MISTA)
    curtas = [c.value for c in at.caption if "= saldo no Domínio" in c.value]
    # a barra impede o markdown de fazer item de lista com o "-" inicial
    assert curtas == [r"\- débitos sem destino R\$ 3.400,00 + crédito sem "
                      r"destino R\$ 500,00 = saldo no Domínio R\$ 5.100,00"]


def test_conta_que_fecha_nao_ganha_ponte_curta():
    at = _tela(GOLDEN)
    assert not any("= saldo no Domínio" in c.value for c in at.caption)


# --- 2. um nome so' pra diferenca a maior numa NF -----------------------------

def test_nf_com_alerta_de_juros_ganha_a_marca_possivel_juros():
    res = _dois_juros()
    linhas = {l["NF"]: l for l in ap.composicao_linhas(
        res.itens, juros=ap.nfs_com_possivel_juros(res.achados))}
    assert linhas["5001 (possível juros)"]["EM ABERTO"] == "-2,46"
    assert linhas["5002 (possível juros)"]["EM ABERTO"] == "-10,00"
    assert not any("pago a maior" in nf for nf in linhas)


def test_linha_do_pago_a_maior_do_lote_continua_com_o_proprio_nome():
    res = _pago_a_maior_e_recebimento_sintetico()
    nomes = [l["NF"] for l in ap.composicao_linhas(
        res.itens, juros=ap.nfs_com_possivel_juros(res.achados))]
    assert "(pago a maior lote 810116)" in nomes


# --- 3. linhas que nao sao NF --------------------------------------------------

def test_linha_que_nao_e_nota_deixa_data_e_compra_vazias_em_vez_de_zero():
    res = _pago_a_maior_e_recebimento_sintetico()
    (linha,) = [l for l in ap.composicao_linhas(res.itens)
                if l["NF"].startswith("(pago a maior")]
    assert linha["Valor compra"] == ""
    assert linha["Data NF"] == "05/12/2025"           # data do lancamento: fica


def test_saldo_anterior_sem_data_fica_vazio_e_nao_travessao():
    res = conferir(_conta([(COMPRA.format(nf="5001"), "0.00", "100.00", 2)],
                          saldo_anterior="-700.00"), CLASSIFICADOR)
    (linha,) = [l for l in ap.composicao_linhas(res.itens)
                if l["NF"] == "(saldo anterior)"]
    assert linha["Data NF"] == ""
    # o trazido fica so' em EM ABERTO, fora do total de compra
    assert linha["Valor compra"] == linha["Total pago"] == ""
    assert linha["EM ABERTO"] == "700,00"


def test_nota_com_zero_continua_mostrando_zero():
    res = _so_credito()
    (linha,) = ap.composicao_linhas(res.itens)
    assert linha["Total pago"] == "0,00"


def test_linhas_sem_nota_marca_so_as_linhas_entre_parenteses():
    res = conferir(_conta([(COMPRA.format(nf="5001"), "0.00", "100.00", 2),
                           (TRANSFERENCIA, "0.00", "40.00", 3)],
                          saldo_anterior="-700.00"), CLASSIFICADOR)
    tabela = ap.composicao_linhas(res.itens)
    marcas = dict(zip((l["NF"] for l in tabela), ap.linhas_sem_nota(res.itens)))
    assert marcas["5001"] is False
    assert marcas["(saldo anterior)"] is True
    assert any(m for nf, m in marcas.items() if nf.startswith("(transfer"))


# --- 5. ponte sem "balde"; nota do saldo anterior em uma frase -----------------

def test_ponte_da_tela_fala_sem_destino_e_nao_balde():
    rotulos = [l["Linha"] for l in ap.ponte_linhas_tela(_so_credito())]
    assert "(+) Créditos sem destino" in rotulos
    rotulos = [l["Linha"] for l in ap.ponte_linhas_tela(_so_debito())]
    assert "(-) Débitos sem destino" in rotulos
    assert not any("balde" in r for r in rotulos)


def test_nota_do_saldo_anterior_consumido_e_uma_frase_so():
    from parser.excel import ExcelParser
    golden = conferir(ExcelParser().parse(GOLDEN)[0], CLASSIFICADOR)
    nota = ap.nota_fora_da_tabela(golden.totais_tabela)
    assert nota == ("O saldo anterior de R$ 429,26 foi pago no período: a "
                    "linha zerou e fica fora dos totais de Valor compra e "
                    "Total pago da tabela.")
    assert nota.count(". ") == 0


# --- 6. icone de Divergencias neutro quando nao ha ----------------------------

def test_divergencias_sem_x_vermelho_quando_nenhuma():
    at = AppTest.from_file(str(MAIN), default_timeout=60).run()
    at.file_uploader[0].set_value(
        ("golden_sintetico.xls", GOLDEN.read_bytes(), "application/vnd.ms-excel")).run()
    titulos = [m.value for m in at.markdown if "Divergências" in m.value]
    assert titulos and not any("❌" in t for t in titulos)


def test_x_vermelho_so_com_divergencia():
    assert "❌" in ap.titulo_divergencias(True)
    assert "❌" not in ap.titulo_divergencias(False)
    assert ap.titulo_divergencias(False).endswith(" Divergências")


# --- 7. juros agrupado nas Observacoes -----------------------------------------

def test_dois_juros_viram_uma_observacao_com_a_lista():
    obs = ap.observacoes_tela(_dois_juros().achados)
    juros = [o for o in obs if "possível juros" in o.frase.lower()]
    assert len(juros) == 1
    (o,) = juros
    assert o.frase.startswith("Possível juros por atraso em 2 notas")
    assert o.frase.count("despesa em conta específica") == 1
    assert o.lista == ("NF 5001: diferença R$ 2,46",
                       "NF 5002: diferença R$ 10,00")


def test_um_juros_so_continua_com_a_frase_inteira():
    res = conferir(_conta([(COMPRA.format(nf="5001"), "0.00", "100.00", 2),
                           (PGTO_NF.format(nf="5001"), "102.46", "0.00", 5)]),
                   CLASSIFICADOR)
    # (a conta fica devedora: a frase dela tambem fala de juros)
    (o,) = [o for o in ap.observacoes_tela(res.achados)
            if "possível juros" in o.frase]
    assert o.lista == ()
    assert "NF 5001" in o.frase and "despesa em conta específica" in o.frase


def test_observacoes_sem_juros_ficam_como_estavam():
    obs = ap.observacoes_tela(_pago_a_maior_e_recebimento_sintetico().achados)
    assert all(o.lista == () for o in obs)


# --- 8. recebimento: so' o fato ------------------------------------------------

def test_recebimento_diz_so_o_fato():
    (item,) = _pago_a_maior_e_recebimento_sintetico().balde
    texto = ap.motivo_humano(item)
    assert "dinheiro recebido do fornecedor, sem nota" in texto.lower()
    assert "mercadoria devolvida" not in texto
    assert "devolução" not in texto


# --- a secao tem o nome da planilha ---------------------------------------------

def test_secao_da_tela_se_chama_itens_sem_destino():
    titulos = [m.value for m in _tela(MISTA).markdown if m.value.startswith("#### ")]
    assert "#### 🪣 Itens sem destino" in titulos
    assert not any("Balde" in t for t in titulos)


# --- compra e pago vazios nas linhas sem nota, total so' das notas -------------

def test_pago_a_maior_do_lote_mostra_valor_so_em_em_aberto():
    res = _pago_a_maior_e_recebimento_sintetico()
    (linha,) = [l for l in ap.composicao_linhas(res.itens)
                if l["NF"].startswith("(pago a maior")]
    assert linha["Valor compra"] == linha["Total pago"] == ""
    assert linha["EM ABERTO"] == "-970,95"


def test_compra_nova_a_mao_com_rotulo_padrao_continua_nota():
    """"(manual: lote X)" tem parenteses no nome, mas e' compra: nao
    fica cinza e mostra a compra."""
    from core.modelo import AlocacaoManual, DestinoAlocacao
    from tests.test_ponte import ESTRANHO
    conta = _conta([(COMPRA.format(nf="5001"), "0.00", "300.00", 2),
                    (ESTRANHO, "0.00", "80.00", 6)])
    item = conferir(conta, CLASSIFICADOR).balde[0].identificador
    res = conferir(conta, CLASSIFICADOR, [AlocacaoManual(
        item_id=item, destino=DestinoAlocacao.COMPRA_NOVA)])
    tabela = ap.composicao_linhas(res.itens)
    marcas = dict(zip((l["NF"] for l in tabela), ap.linhas_sem_nota(res.itens)))
    (nf,) = [n for n in marcas if n.startswith("(manual")]
    assert marcas[nf] is False
    assert {l["NF"]: l for l in tabela}[nf]["Valor compra"] == "80,00"


def test_tooltip_do_total_diz_que_compra_e_pago_sao_so_das_notas():
    for rotulo in (ap.ROTULO_TOTAL_EM_ABERTO, ap.ROTULO_TOTAL_COM_QUITADAS):
        assert "somam só as notas" in ap.ajuda_total(rotulo)


def test_planilha_deixa_compra_e_pago_vazios_e_explica_o_total():
    import io
    import openpyxl
    from app.pipeline import processar_excel
    from relatorio.excel import ABA_CONFERENCIA, gerar_excel
    res = processar_excel(MISTA).resultados
    ws = openpyxl.load_workbook(io.BytesIO(gerar_excel(
        res, MISTA.read_bytes(), MISTA.name)))[ABA_CONFERENCIA]
    cab = next(c.row for linha in ws.iter_rows() for c in linha if c.value == "NF")
    nomes = {c.value: c.column for c in ws[cab] if c.value}
    linha = next(c.row for c in ws["A"] if isinstance(c.value, str)
                 and c.value.startswith("(transfer"))
    assert ws.cell(linha, nomes["Valor compra"]).value is None
    assert ws.cell(linha, nomes["Total pago"]).value is None
    assert ws.cell(linha, nomes["EM ABERTO"]).value is not None
    textos = [c.value for l in ws.iter_rows() for c in l if isinstance(c.value, str)]
    assert any("somam só as notas" in t for t in textos)
