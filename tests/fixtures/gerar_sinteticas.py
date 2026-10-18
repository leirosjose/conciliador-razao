"""Gerador deterministico das planilhas sinteticas de teste, no layout do
relatorio Razao do Dominio. Tudo e' ficticio: empresas, fornecedores e
escritorio levam FICTICIA/FICTICIO no nome, o unico CNPJ e' 00.000.000/0000-00,
NFs, lotes, datas e valores sao inventados.

Uso:  python tests/fixtures/gerar_sinteticas.py [DESTINO]
      (DESTINO padrao: a pasta deste arquivo)
Gera:
  sintetica_integra.xls, sintetica_corrompida.xls, sintetica_adulterada.xls
  sintetica_integra.xlsx
  golden_sintetico.xls
  cenario_1_conciliado.xls ... cenario_6_pago_a_maior.xls
Rodar duas vezes gera os mesmos bytes.

SINTETICAS (balde, alocacao, reparo do .xls). Conta 7777 FORNECEDOR FICTICIO
LTDA, saldo anterior -10.000,00, saldo na COLUNA 11 (exercita a deteccao pelo
cabecalho; o golden usa a 12). Cobre: compra multi-linha da mesma NF,
pagamento por numero, parcelas iguais, retencao com NF, baixa de adiantamento
== |saldo anterior| exato, baixa parcial (!= agregado -> balde), multi-NF com
soma exata (casa um-pra-muitos), estorno, credito nao classificado, debito
sem padrao. Identidade: soma do aberto 8.000,00 (1.900 + 750 + 250 + 5.100; os
250,00 sao a transferencia de saldos, que entra na composicao); alvo 5.100,00;
gap 2.900,00 == soma(D-C) do balde (D 3.400,00 / C 500,00).

A corrompida reproduz o defeito do exportador do Dominio e e' o INVERSO do
reparo do parser (parser/excel.py): records BLANK (0x0201) orfaos inseridos
logo apos o EOF (0x000A) dos globals SEM atualizar o offset do BOUNDSHEET
(0x0085) -- o ponteiro continua apontando pra posicao antiga do BOF da sheet,
que agora e' o meio do bloco de BLANKs. O xlrd puro estoura com
"Expected BOF record; found b'\\x01\\x02...'". A adulterada e' a corrompida
com um debito reescrito (10.000 -> 9.000): repara, mas o self-check de totais
nao fecha -> o parser rejeita.

A .xlsx tem o MESMO conteudo da integra, com os tracos de um export .xlsx do
Dominio: cabecalho 'Lote' no lugar de 'Número', datas como datetime nativo,
dinheiro float/int, saldo na coluna 11 com sinal negativo. Exercita a rota
openpyxl com a mesma identidade.

GOLDEN SINTETICO. 12 lancamentos; o saldo anterior (-429,26) e' quitado por 2
parcelas de uma NF antiga; 2 NFs em aberto com 4 pagamentos parciais cada,
uma parcela 0,18 abaixo das outras. Folha, CONSOLIDADO, cabecalho na linha 8,
saldo na coluna 12, Total da conta, rodape do sistema e composicao manual nas
ultimas linhas (com o artefato de float de conta feita no Excel). O
fornecedor aparece como "SA" no cabecalho da conta e nas compras e como
"S.A." nos pagamentos. Esperado: FECHA; NF 618204 aberto 2.820,00, NF 274591
aberto 1.440,09, total 4.260,09 = |saldo final|; balde vazio.

CENARIOS. Uma razao por cenario, NFs 7xxxxx-9xxxxx, lotes 31xxxxx; NFs, lotes
e valores de um cenario nao se repetem em outro, e o saldo corrente nunca
zera (conferir_unicidade):
  1 conciliado    tudo casa por numero; uma NF quitada com IRRF
  2 centavo       3 parcelas por valor, sem numero; quitada com arredondamento de 0,01
  3 sem par       boleto sem NF e sem valor de parcela -> balde; INCOMPLETO
  4 estorno       baixa de adiantamento estornada -> par anulado
  5 duplicidade   a mesma NF lancada duas vezes -> alerta de compra repetida
  6 pago a maior  pagamento cita 2 NFs e passa da soma -> linha "pago a maior"
"""
from __future__ import annotations

import re
import struct
import sys
import zipfile
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
import xlwt
from xlrd import xldate
from xlwt.CompoundDoc import XlsDoc

DESTINO = Path(__file__).resolve().parent

D = Decimal
CNPJ = "00.000.000/0000-00"
RODAPE = "Sistema licenciado para ESCRITORIO FICTICIO LTDA"
TIMESTAMP_FIXO = datetime(2026, 1, 1)     # carimbo fixo da .xlsx: regravar nao muda bytes


def _serial(d: date) -> float:
    return xldate.xldate_from_date_tuple((d.year, d.month, d.day), 0)


# === SINTETICAS ==============================================================

EMPRESA = "COMERCIAL FICTICIA LTDA"
PERIODO = "01/01/2024 - 31/12/2025"
CONTA_CODIGO = 7777
CONTA_CLASSIFICACAO = "2.01.01.001.00"
CONTA_NOME = "FORNECEDOR FICTICIO LTDA"
SALDO_ANTERIOR = Decimal("-10000.00")
COL_SALDO = 11                # coluna MOVEL: golden traz na 12, esta na 11

# (data, lote, historico, debito, credito). A ORDEM faz o saldo-exercicio
# NUNCA zerar: zerar ativaria o corte por saldo zero e mandaria o balde
# inteiro pro periodo zerado. O corte tem sinteticas proprias (tests/test_corte.py).
LANCAMENTOS = [
    (date(2024, 1, 15), 3002, "Compras conf. NF Vr. compras conf. 1111 FORNECEDOR FICTICIO LTDA",
     None, Decimal("1000.00")),
    (date(2024, 1, 15), 3003, "Vr. compras conf. 1111 FORNECEDOR FICTICIO LTDA",
     None, Decimal("2100.00")),
    (date(2024, 2, 1), 3001, "Valor ref. a BAIXA DE ADIANTAMENTO A FORN - FORNECEDOR FICTICIO",
     Decimal("10000.00"), None),
    (date(2024, 4, 1), 3004, "Valor ref. a PGTO A FORNECEDOR FICTICIO NF 1111",
     Decimal("1200.00"), None),
    (date(2024, 5, 1), 3005, "Compras conf. NF Vr. compras conf. 2222 FORNECEDOR FICTICIO LTDA",
     None, Decimal("2400.00")),
    (date(2024, 6, 3), 3006, "Valor ref. a Pagamento de parcela a fornecedor",
     Decimal("1200.00"), None),
    (date(2024, 7, 1), 3007, "Valor ref. a Pagamento de parcela a fornecedor",
     Decimal("1200.00"), None),
    (date(2024, 8, 1), 3008, "Vr. compras conf. 3333 FORNECEDOR FICTICIO LTDA",
     None, Decimal("900.00")),
    (date(2024, 9, 2), 3009, "IRRF RETIDO S/ NF 3333",
     Decimal("150.00"), None),
    (date(2024, 10, 1), 3010, "Compras conf. NF Vr. compras conf. 4444 FORNECEDOR FICTICIO LTDA",
     None, Decimal("650.00")),
    (date(2024, 10, 1), 3011, "Vr. compras conf. 5555 FORNECEDOR FICTICIO LTDA",
     None, Decimal("350.00")),
    (date(2024, 11, 1), 3012, "Valor ref. a PGTO A FORNECEDOR FICTICIO NF 4444 E NFS: 5555",
     Decimal("1000.00"), None),
    (date(2025, 1, 15), 3013, "Valor ref. a BAIXA DE ADIANTAMENTO A FORN - FORNECEDOR FICTICIO",
     Decimal("3000.00"), None),
    (date(2025, 2, 3), 3014, "Valor ref. estorno de BAIXA DE ADIANTAMENTO indevida - FORNECEDOR FICTICIO",
     None, Decimal("500.00")),
    (date(2025, 3, 3), 3015, "Valor ref. a TRANSFERENCIA DE SALDOS DE CONTAS DE UM MESMO FORNECEDOR",
     None, Decimal("250.00")),
    (date(2025, 3, 20), 3017, "Vr. compras conf. 6666 FORNECEDOR FICTICIO LTDA",
     None, Decimal("5100.00")),
    (date(2025, 4, 1), 3016, "TARIFA ADMINISTRATIVA SEM PADRAO CONHECIDO",
     Decimal("400.00"), None),
]

LINHA_PRIMEIRO_LANCAMENTO = 8          # linha (0-based) do 1o lancamento na sheet
COL_DEBITO = 8

# records BIFF usados na cirurgia
BIFF_EOF = 0x000A
BIFF_BLANK = 0x0201
BIFF_RK = 0x027E
BIFF_NUMBER = 0x0203
BIFF_MULRK = 0x00BD           # celulas numericas adjacentes fundidas pelo xlwt


def montar_stream_biff() -> bytes:
    """Monta o .xls integro no contrato do Dominio e devolve o stream BIFF cru."""
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Razão")

    ws.write(0, 0, "Empresa:");  ws.write(0, 2, EMPRESA)
    ws.write(1, 0, "C.N.P.J.:"); ws.write(1, 2, CNPJ)
    ws.write(2, 0, "Período:");  ws.write(2, 2, PERIODO)
    ws.write(3, 0, "RAZÃO")
    # linha 4 vazia
    ws.write(5, 0, "Data");        ws.write(5, 1, "Número")
    ws.write(5, 2, "Histórico");   ws.write(5, 7, "Cta.C.Part.")
    ws.write(5, 8, "Débito");      ws.write(5, 9, "Crédito")
    ws.write(5, COL_SALDO, "Saldo-Exercício")
    ws.write(6, 0, "Conta:");      ws.write(6, 1, float(CONTA_CODIGO))
    ws.write(6, 2, CONTA_CLASSIFICACAO); ws.write(6, 5, CONTA_NOME)
    ws.write(7, 2, "SALDO ANTERIOR");    ws.write(7, COL_SALDO, float(SALDO_ANTERIOR))

    saldo = SALDO_ANTERIOR
    total_d = total_c = Decimal("0.00")
    for i, (dia, lote, hist, deb, cred) in enumerate(LANCAMENTOS):
        r = LINHA_PRIMEIRO_LANCAMENTO + i
        ws.write(r, 0, _serial(dia))
        ws.write(r, 1, float(lote))
        ws.write(r, 2, hist)
        ws.write(r, 7, 38.0)                       # Cta.C.Part. generica
        if deb is not None:
            ws.write(r, COL_DEBITO, float(deb))
            saldo += deb
            total_d += deb
        if cred is not None:
            ws.write(r, 9, float(cred))
            saldo -= cred
            total_c += cred
        ws.write(r, COL_SALDO, float(saldo))       # saldo corrente linha a linha

    r_total = LINHA_PRIMEIRO_LANCAMENTO + len(LANCAMENTOS) + 1
    ws.write(r_total, 4, "Total da conta:")
    ws.write(r_total, 8, float(total_d))
    ws.write(r_total, 9, float(total_c))
    return wb.get_biff_data()


def corromper(stream: bytes, n_blanks: int = 30) -> bytes:
    """Insere BLANKs orfaos apos o EOF dos globals SEM atualizar o offset do
    BOUNDSHEET -- o inverso exato do reparo do parser."""
    pos = 0
    while pos + 4 <= len(stream):
        rec_id = int.from_bytes(stream[pos:pos + 2], "little")
        rec_len = int.from_bytes(stream[pos + 2:pos + 4], "little")
        fim = pos + 4 + rec_len
        if rec_id == BIFF_EOF:                     # primeiro EOF = fim dos globals
            blank = b"".join(
                BIFF_BLANK.to_bytes(2, "little") + (6).to_bytes(2, "little")
                + i.to_bytes(2, "little") + (2).to_bytes(2, "little")
                + (0).to_bytes(2, "little")
                for i in range(n_blanks))
            return stream[:fim] + blank + stream[fim:]
        pos = fim
    raise ValueError("EOF dos globals nao encontrado no stream BIFF")


def adulterar_celula(stream: bytes, linha: int, coluna: int, novo_inteiro: int) -> bytes:
    """Reescreve o valor de uma celula numerica (RK ou NUMBER) -- produz o
    arquivo que repara mas NAO passa no self-check (teste do gate)."""
    data = bytearray(stream)
    pos = 0
    while pos + 4 <= len(data):
        rec_id = int.from_bytes(data[pos:pos + 2], "little")
        rec_len = int.from_bytes(data[pos + 2:pos + 4], "little")
        if rec_id in (BIFF_RK, BIFF_NUMBER):
            r = int.from_bytes(data[pos + 4:pos + 6], "little")
            c = int.from_bytes(data[pos + 6:pos + 8], "little")
            if r == linha and c == coluna:
                if rec_id == BIFF_RK:              # rk int: (v << 2) | 0b10
                    data[pos + 10:pos + 14] = ((novo_inteiro << 2) | 0b10).to_bytes(4, "little")
                else:                              # NUMBER: double em pos+10
                    data[pos + 10:pos + 18] = struct.pack("<d", float(novo_inteiro))
                return bytes(data)
        elif rec_id == BIFF_MULRK:
            # MULRK: row(2) colFirst(2) [xf(2) rk(4)]*N colLast(2)
            r = int.from_bytes(data[pos + 4:pos + 6], "little")
            col_ini = int.from_bytes(data[pos + 6:pos + 8], "little")
            col_fim = int.from_bytes(data[pos + 4 + rec_len - 2:pos + 4 + rec_len], "little")
            if r == linha and col_ini <= coluna <= col_fim:
                rk_pos = pos + 8 + (coluna - col_ini) * 6 + 2
                data[rk_pos:rk_pos + 4] = ((novo_inteiro << 2) | 0b10).to_bytes(4, "little")
                return bytes(data)
        pos += 4 + rec_len
    raise ValueError(f"celula numerica ({linha},{coluna}) nao encontrada")


def montar_xlsx() -> bytes:
    """O MESMO conteudo da BIFF integra, no formato .xlsx (ver o docstring)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Razão"

    ws.cell(1, 1, "Empresa:");  ws.cell(1, 3, EMPRESA)
    ws.cell(2, 1, "C.N.P.J.:"); ws.cell(2, 3, CNPJ)
    ws.cell(3, 1, "Período:");  ws.cell(3, 3, PERIODO)
    ws.cell(4, 1, "RAZÃO")
    ws.cell(6, 1, "Data");        ws.cell(6, 2, "Lote")
    ws.cell(6, 3, "Histórico");   ws.cell(6, 8, "Cta.C.Part.")
    ws.cell(6, 9, "Débito");      ws.cell(6, 10, "Crédito")
    ws.cell(6, COL_SALDO + 1, "Saldo-Exercício")
    ws.cell(7, 1, "Conta:");      ws.cell(7, 2, CONTA_CODIGO)
    ws.cell(7, 3, CONTA_CLASSIFICACAO); ws.cell(7, 6, CONTA_NOME)
    ws.cell(8, 3, "SALDO ANTERIOR");    ws.cell(8, COL_SALDO + 1, float(SALDO_ANTERIOR))

    saldo = SALDO_ANTERIOR
    total_d = total_c = Decimal("0.00")
    for i, (dia, lote, hist, deb, cred) in enumerate(LANCAMENTOS):
        r = LINHA_PRIMEIRO_LANCAMENTO + 1 + i          # openpyxl e' 1-based
        ws.cell(r, 1, datetime(dia.year, dia.month, dia.day))
        ws.cell(r, 2, lote)
        ws.cell(r, 3, hist)
        ws.cell(r, 8, 38)
        if deb is not None:
            ws.cell(r, COL_DEBITO + 1, float(deb))
            saldo += deb
            total_d += deb
        if cred is not None:
            ws.cell(r, 10, float(cred))
            saldo -= cred
            total_c += cred
        ws.cell(r, COL_SALDO + 1, float(saldo))

    r_total = LINHA_PRIMEIRO_LANCAMENTO + 1 + len(LANCAMENTOS) + 1
    ws.cell(r_total, 5, "Total da conta:")
    ws.cell(r_total, 9, float(total_d))
    ws.cell(r_total, 10, float(total_c))

    wb.properties.created = TIMESTAMP_FIXO
    wb.properties.modified = TIMESTAMP_FIXO
    buf = BytesIO()
    wb.save(buf)
    return _zip_deterministico(buf.getvalue())


_MODIFICADO = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _zip_deterministico(data: bytes) -> bytes:
    """Regrava o ZIP com timestamps fixos e entradas ordenadas. O openpyxl
    carimba a hora da maquina em cada entrada do ZIP e, no save, no
    dcterms:modified do docProps/core.xml; os dois voltam pro TIMESTAMP_FIXO.
    O conteudo das celulas nao muda."""
    fixo = TIMESTAMP_FIXO.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    origem = zipfile.ZipFile(BytesIO(data))
    saida = BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z:
        for nome in sorted(origem.namelist()):
            conteudo = origem.read(nome)
            if nome == "docProps/core.xml":
                conteudo, n = _MODIFICADO.subn(rb"\g<1>" + fixo + rb"\g<2>", conteudo)
                assert n == 1, "dcterms:modified nao encontrado no core.xml"
            info = zipfile.ZipInfo(nome, date_time=TIMESTAMP_FIXO.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, conteudo)
    return saida.getvalue()


# === GOLDEN SINTETICO ========================================================

G_EMPRESA = "INDUSTRIA FICTICIA DE ALIMENTOS LTDA"
G_PERIODO = "01/01/2024 - 30/04/2025"
G_CONSOLIDADO = "CONSOLIDADO (Empresas: 101,102,103,104)"
G_CONTA_CODIGO = 5218
G_CONTA_CLASSIFICACAO = "2.01.01.03.500"
G_FORN_PGTO = "FERRAMENTAS FICTICIAS COMERCIO S.A."
G_FORN_COMPRA = "FERRAMENTAS FICTICIAS COMERCIO SA"
G_SALDO_ANTERIOR = Decimal("-429.26")
G_COL_SALDO = 12
G_DATA = xlwt.easyxf(num_format_str="m/d/yy")

NF_ANTIGA, NF_A, NF_B = "512309", "618204", "274591"


def _g_pg(nf, parc):
    return f"- PAGTO REF. NF.:  {nf}.{parc} - PARC.:{G_FORN_PGTO}"


def _g_compra(nf):
    return f"Compras conf. NF Vr. compras conf. {nf} {G_FORN_COMPRA}"


# (data, lote, historico, cta_part, debito, credito)
G_LANCAMENTOS = [
    (date(2024, 1, 20), 2310457, _g_pg(NF_ANTIGA, 4), 7311, D("214.63"), None),
    (date(2024, 2, 20), 2329718, _g_pg(NF_ANTIGA, 5), 7311, D("214.63"), None),
    (date(2024, 12, 11), 2518642, _g_compra(NF_A), 2208, None, D("4700.00")),
    (date(2025, 1, 2), 2529370, _g_compra(NF_B), 2208, None, D("2400.03")),
    (date(2025, 1, 9), 2526151, _g_pg(NF_A, 1), 7340, D("470.00"), None),
    (date(2025, 1, 9), 2526483, _g_pg(NF_B, 1), 7340, D("240.03"), None),
    (date(2025, 2, 9), 2540512, _g_pg(NF_A, 2), 7340, D("470.00"), None),
    (date(2025, 2, 9), 2540586, _g_pg(NF_B, 2), 7340, D("239.85"), None),
    (date(2025, 3, 9), 2558273, _g_pg(NF_A, 3), 7340, D("470.00"), None),
    (date(2025, 3, 9), 2558319, _g_pg(NF_B, 3), 7340, D("240.03"), None),
    (date(2025, 4, 9), 2572604, _g_pg(NF_A, 4), 7340, D("470.00"), None),
    (date(2025, 4, 9), 2572751, _g_pg(NF_B, 4), 7340, D("240.03"), None),
]
# a composicao manual traz o aberto da NF B com artefato de float (conta feita no Excel)
ARTEFATO_NF_B = 1440.090000000001          # 16 algarismos: o openpyxl grava com %.16g
COMPOSICAO_MANUAL = [(f"nf {NF_A}", D("2820.00")), (f"nf {NF_B}", D("1440.09"))]


def montar_golden() -> bytes:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Razão")
    ws.write(0, 0, "Empresa:"); ws.write(0, 2, G_EMPRESA)
    ws.write(0, 11, "Folha:"); ws.write(0, 13, 1.0)
    ws.write(1, 0, "C.N.P.J.:"); ws.write(1, 2, CNPJ)
    ws.write(2, 0, "Período:"); ws.write(2, 2, G_PERIODO)
    ws.write(3, 0, G_CONSOLIDADO)
    ws.write(5, 0, "RAZÃO")
    ws.write(7, 0, "Data"); ws.write(7, 1, "Número"); ws.write(7, 2, "Histórico")
    ws.write(7, 7, "Cta.C.Part."); ws.write(7, 8, "Débito"); ws.write(7, 9, "Crédito")
    ws.write(7, G_COL_SALDO, "Saldo-Exercício")
    ws.write(8, 0, "Conta:"); ws.write(8, 1, float(G_CONTA_CODIGO))
    ws.write(8, 2, G_CONTA_CLASSIFICACAO); ws.write(8, 5, G_FORN_COMPRA)
    ws.write(9, 2, "SALDO ANTERIOR"); ws.write(9, G_COL_SALDO, float(G_SALDO_ANTERIOR))
    saldo, td, tc = G_SALDO_ANTERIOR, D("0"), D("0")
    for i, (dia, lote, hist, cp, deb, cred) in enumerate(G_LANCAMENTOS):
        r = 10 + i
        ws.write(r, 0, _serial(dia), G_DATA); ws.write(r, 1, float(lote)); ws.write(r, 2, hist)
        ws.write(r, 7, float(cp))
        if deb is not None:
            ws.write(r, 8, float(deb)); saldo += deb; td += deb
        if cred is not None:
            ws.write(r, 9, float(cred)); saldo -= cred; tc += cred
        ws.write(r, G_COL_SALDO, float(saldo))
    r = 10 + len(G_LANCAMENTOS) + 1
    ws.write(r, 4, "Total da conta:"); ws.write(r, 8, float(td)); ws.write(r, 9, float(tc))
    ws.write(r + 2, 0, RODAPE)
    ws.write(r + 3, 9, "composição")
    for k, (rot, v) in enumerate(COMPOSICAO_MANUAL):
        ws.write(r + 4 + k, 9, rot); ws.write(r + 4 + k, 12, ARTEFATO_NF_B if k == 1 else float(v))
    ws.write(r + 4 + len(COMPOSICAO_MANUAL), 12, float(sum(v for _, v in COMPOSICAO_MANUAL)))
    return wb.get_biff_data()


# === CENARIOS ================================================================

C_COL_SALDO = 12
C_LINHA_LANC = 10            # linha (0-based) do 1o lancamento


# os tres formatos de historico que o classificador reconhece
def pg(nf: str, parc: int, forn: str) -> str:
    """Pagamento que CITA a NF (casa por numero)."""
    return f"- PAGTO REF. NF.:  {nf}.{parc} - PARC.:{forn}"


def compra(nf: str, forn: str) -> str:
    return f"Compras conf. NF Vr. compras conf. {nf} {forn}"


def irrf(nf: str) -> str:
    return f"IRRF RETIDO S/ NF {nf}"


# (data, lote, historico, cta_part, debito, credito)
CENARIOS: dict[str, dict] = {
    # 1. tudo casa por numero; uma NF quitada (com IRRF), uma parcelada em
    #    aberto; o saldo anterior e' quitado por pagamento da NF antiga.
    "cenario_1_conciliado": dict(
        empresa="COMERCIO FICTICIO DE EMBALAGENS LTDA",
        periodo="01/01/2025 - 31/03/2025",
        conta=5301, classificacao="2.01.01.03.611",
        fornecedor="DISTRIBUIDORA FICTICIA DE PAPEIS LTDA",
        saldo_anterior=D("-1873.46"),
        lancamentos=lambda f: [
            (date(2025, 1, 6), 3102958, compra("791306", f), 2215, None, D("4316.80")),
            (date(2025, 1, 10), 3104417, pg("784512", 3, f), 7362, D("1873.46"), None),
            (date(2025, 2, 3), 3118630, compra("793018", f), 2215, None, D("2957.35")),
            (date(2025, 2, 10), 3121574, pg("791306", 1, f), 7362, D("1438.93"), None),
            (date(2025, 3, 5), 3135209, irrf("793018"), 4421, D("44.36"), None),
            (date(2025, 3, 5), 3135242, pg("793018", 1, f), 7362, D("2912.99"), None),
            (date(2025, 3, 10), 3137761, pg("791306", 2, f), 7362, D("1438.93"), None),
        ],
    ),
    # 2. NF de 2.853,47 paga em 3 x 951,16 SEM numero no historico: casa por
    #    valor (1 de N parcelas, tolerancia de 1 centavo por parcela) e a soma
    #    passa 0,01 da nota -> "quitada com arredondamento".
    "cenario_2_centavo": dict(
        empresa="INDUSTRIA FICTICIA DE MOVEIS LTDA",
        periodo="01/01/2025 - 30/04/2025",
        conta=5302, classificacao="2.01.01.03.612",
        fornecedor="MADEIREIRA FICTICIA DO VALE LTDA",
        saldo_anterior=D("0.00"),
        lancamentos=lambda f: [
            (date(2025, 1, 14), 3106182, compra("802741", f), 2231, None, D("2853.47")),
            (date(2025, 2, 14), 3123905, f"PAGAMENTO DE DUPLICATA - {f}", 7385, D("951.16"), None),
            (date(2025, 3, 14), 3139458, f"PAGAMENTO DE DUPLICATA - {f}", 7385, D("951.16"), None),
            (date(2025, 3, 20), 3141027, compra("805396", f), 2231, None, D("1586.90")),
            (date(2025, 4, 14), 3152614, f"PAGAMENTO DE DUPLICATA - {f}", 7385, D("951.16"), None),
        ],
    ),
    # 3. boleto pago sem NF e com valor que nao e' parcela de nota nenhuma ->
    #    balde SEM_VINCULO; a conta fica INCOMPLETA (o balde explica o gap).
    "cenario_3_sem_par": dict(
        empresa="TRANSPORTADORA FICTICIA LTDA",
        periodo="01/01/2025 - 30/04/2025",
        conta=5303, classificacao="2.01.01.03.613",
        fornecedor="AUTOPECAS FICTICIAS DO NORTE LTDA",
        saldo_anterior=D("-3512.84"),
        lancamentos=lambda f: [
            (date(2025, 1, 8), 3103376, pg("781945", 3, f), 7394, D("1756.42"), None),
            (date(2025, 1, 22), 3110849, compra("886120", f), 2246, None, D("5827.40")),
            (date(2025, 2, 7), 3117935, pg("781945", 4, f), 7394, D("1756.42"), None),
            (date(2025, 2, 24), 3126518, compra("889457", f), 2246, None, D("3419.65")),
            (date(2025, 3, 7), 3133602, pg("886120", 1, f), 7394, D("2913.70"), None),
            (date(2025, 3, 18), 3138814,
             "PGTO BOLETO AUTOPECAS FICTICIAS DO NORTE - REF. ACORDO COMERCIAL",
             7394, D("1264.38"), None),
            (date(2025, 4, 7), 3149027, pg("886120", 2, f), 7394, D("2913.70"), None),
        ],
    ),
    # 4. baixa de adiantamento lancada por engano e estornada 15 dias depois:
    #    mesmo valor, mesma contrapartida, o estorno DECLARA o que reverte ->
    #    par anulado, fora do balde e da composicao.
    "cenario_4_estorno": dict(
        empresa="SUPERMERCADO FICTICIO LTDA",
        periodo="01/01/2025 - 30/04/2025",
        conta=5304, classificacao="2.01.01.03.614",
        fornecedor="LATICINIOS FICTICIOS SERRANOS LTDA",
        saldo_anterior=D("-2236.71"),
        lancamentos=lambda f: [
            (date(2025, 1, 9), 3104802, compra("912604", f), 2253, None, D("3684.27")),
            (date(2025, 1, 15), 3107391, pg("907318", 2, f), 7418, D("2236.71"), None),
            (date(2025, 2, 4), 3118057,
             "Valor ref. a BAIXA DE ADIANTAMENTO A FORN - LATICINIOS FICTICIOS SERRANOS",
             1185, D("1947.52"), None),
            (date(2025, 2, 11), 3121993, compra("915270", f), 2253, None, D("2608.09")),
            (date(2025, 2, 19), 3125648,
             "Valor ref. estorno de BAIXA DE ADIANTAMENTO indevida - LATICINIOS FICTICIOS SERRANOS",
             1185, None, D("1947.52")),
            (date(2025, 3, 10), 3135901, pg("912604", 1, f), 7418, D("3684.27"), None),
            (date(2025, 4, 10), 3150216, pg("915270", 1, f), 7418, D("1304.05"), None),
        ],
    ),
    # 5. a mesma NF lancada duas vezes (lotes e datas diferentes, mesmo valor)
    #    e paga uma vez: alerta de compra repetida; a composicao NAO tira a
    #    copia -- mostra a NF com o dobro e metade em aberto, como a razao.
    "cenario_5_duplicidade": dict(
        empresa="RESTAURANTE FICTICIO LTDA",
        periodo="01/01/2025 - 31/03/2025",
        conta=5305, classificacao="2.01.01.03.615",
        fornecedor="HORTIFRUTI FICTICIO DO CAMPO LTDA",
        saldo_anterior=D("-1392.58"),
        lancamentos=lambda f: [
            (date(2025, 1, 7), 3102671, compra("931582", f), 2267, None, D("2163.94")),
            (date(2025, 1, 13), 3105739, pg("928140", 1, f), 7425, D("1392.58"), None),
            (date(2025, 1, 20), 3109418, compra("931582", f), 2267, None, D("2163.94")),
            (date(2025, 2, 7), 3117264, pg("931582", 1, f), 7425, D("2163.94"), None),
            (date(2025, 2, 18), 3123087, compra("934027", f), 2267, None, D("1478.61")),
            (date(2025, 3, 18), 3138526, pg("934027", 1, f), 7425, D("1478.61"), None),
        ],
    ),
    # 6. um pagamento cita duas NFs e paga 177,38 acima da soma delas: cada
    #    NF citada quita, a sobra vira a linha "(pago a maior lote X)".
    "cenario_6_pago_a_maior": dict(
        empresa="CONSTRUTORA FICTICIA LTDA",
        periodo="01/01/2025 - 30/04/2025",
        conta=5306, classificacao="2.01.01.03.616",
        fornecedor="MATERIAIS FICTICIOS DE CONSTRUCAO LTDA",
        saldo_anterior=D("0.00"),
        lancamentos=lambda f: [
            (date(2025, 2, 3), 3115820, compra("947215", f), 2272, None, D("1837.26")),
            (date(2025, 2, 3), 3115847, compra("947216", f), 2272, None, D("966.43")),
            (date(2025, 2, 20), 3126104, compra("952608", f), 2272, None, D("4105.72")),
            (date(2025, 3, 5), 3133419,
             "Valor ref. a PGTO A MATERIAIS FICTICIOS NFS:947215 E NF 947216",
             7431, D("2981.07"), None),
            (date(2025, 4, 22), 3155290, pg("952608", 1, f), 7431, D("2052.86"), None),
        ],
    ),
}


def lancamentos(cen: dict) -> list[tuple]:
    return cen["lancamentos"](cen["fornecedor"])


def montar_cenario(cen: dict) -> bytes:
    """Uma conta, no layout do relatorio Razao do Dominio, com saldo corrente
    linha a linha, 'Total da conta' e rodape do sistema."""
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Razão")
    ws.write(0, 0, "Empresa:"); ws.write(0, 2, cen["empresa"])
    ws.write(0, 11, "Folha:"); ws.write(0, 13, 1.0)
    ws.write(1, 0, "C.N.P.J.:"); ws.write(1, 2, CNPJ)
    ws.write(2, 0, "Período:"); ws.write(2, 2, cen["periodo"])
    ws.write(5, 0, "RAZÃO")
    ws.write(7, 0, "Data"); ws.write(7, 1, "Número"); ws.write(7, 2, "Histórico")
    ws.write(7, 7, "Cta.C.Part."); ws.write(7, 8, "Débito"); ws.write(7, 9, "Crédito")
    ws.write(7, C_COL_SALDO, "Saldo-Exercício")
    ws.write(8, 0, "Conta:"); ws.write(8, 1, float(cen["conta"]))
    ws.write(8, 2, cen["classificacao"]); ws.write(8, 5, cen["fornecedor"])
    ws.write(9, 2, "SALDO ANTERIOR"); ws.write(9, C_COL_SALDO, float(cen["saldo_anterior"]))
    saldo, td, tc = cen["saldo_anterior"], D("0.00"), D("0.00")
    for i, (dia, lote, hist, cp, deb, cred) in enumerate(lancamentos(cen)):
        r = C_LINHA_LANC + i
        ws.write(r, 0, _serial(dia)); ws.write(r, 1, float(lote)); ws.write(r, 2, hist)
        ws.write(r, 7, float(cp))
        if deb is not None:
            ws.write(r, 8, float(deb)); saldo += deb; td += deb
        if cred is not None:
            ws.write(r, 9, float(cred)); saldo -= cred; tc += cred
        ws.write(r, C_COL_SALDO, float(saldo))
    r = C_LINHA_LANC + len(lancamentos(cen)) + 1
    ws.write(r, 4, "Total da conta:"); ws.write(r, 8, float(td)); ws.write(r, 9, float(tc))
    ws.write(r + 2, 0, RODAPE)
    return wb.get_biff_data()


def conferir_unicidade() -> None:
    """NFs, lotes e valores de um cenario nao se repetem em outro (dentro do
    mesmo cenario a repeticao e' o proprio caso: parcelas iguais, estorno,
    compra repetida). O saldo corrente nunca zera: zerar ativaria o corte por
    saldo zero, que nao e' o que estes exemplos mostram."""
    dono: dict[tuple[str, str], str] = {}
    for nome, cen in CENARIOS.items():
        chaves = {("valor", str(abs(cen["saldo_anterior"])))} - {("valor", "0.00")}
        saldo = cen["saldo_anterior"]
        for _, lote, hist, _, deb, cred in lancamentos(cen):
            chaves.add(("lote", str(lote)))
            chaves.add(("valor", str(deb if deb is not None else cred)))
            chaves.update(("nf", t.split(".")[0]) for t in hist.replace(":", " ").split()
                          if t.split(".")[0].isdigit() and len(t.split(".")[0]) == 6)
            saldo += (deb or D("0")) - (cred or D("0"))
            assert saldo != D("0.00"), f"{nome}: saldo corrente zera no lote {lote}"
        for k in chaves:
            assert dono.setdefault(k, nome) == nome, f"{k} em {dono[k]} e {nome}"


# === GERACAO =================================================================

def gerar(destino: Path = DESTINO) -> None:
    destino.mkdir(parents=True, exist_ok=True)
    integro = montar_stream_biff()
    corrompido = corromper(integro)
    # a baixa de 10.000 -> 9.000, achada pela LINHA DO LOTE (nao e' o 1o lancamento)
    linha_baixa = LINHA_PRIMEIRO_LANCAMENTO + next(
        i for i, l in enumerate(LANCAMENTOS) if l[1] == 3001)
    adulterado = adulterar_celula(corrompido, linha_baixa, COL_DEBITO, 9000)

    for nome, stream in (("sintetica_integra.xls", integro),
                         ("sintetica_corrompida.xls", corrompido),
                         ("sintetica_adulterada.xls", adulterado)):
        XlsDoc().save(str(destino / nome), stream)     # container OLE real
        print(f"gerado: {destino / nome}")

    (destino / "sintetica_integra.xlsx").write_bytes(montar_xlsx())
    print(f"gerado: {destino / 'sintetica_integra.xlsx'}")

    XlsDoc().save(str(destino / "golden_sintetico.xls"), montar_golden())
    print(f"gerado: {destino / 'golden_sintetico.xls'}")

    conferir_unicidade()
    for nome, cen in CENARIOS.items():
        XlsDoc().save(str(destino / f"{nome}.xls"), montar_cenario(cen))
        print(f"gerado: {destino / nome}.xls")


if __name__ == "__main__":
    gerar(Path(sys.argv[1]) if len(sys.argv) > 1 else DESTINO)
