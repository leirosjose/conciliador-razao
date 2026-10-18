"""Textos e fio do cabecalho; somente entradas sinteticas."""
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path
import ast
import re

import pytest
from streamlit.testing.v1 import AppTest
from app import apresentacao as a
from app.pipeline import processar_excel
from core import alocacao
from core.modelo import ComposicaoItem, Origem
from parser.excel import ExcelParser, ErroLeituraExcel

RAIZ = Path(__file__).resolve().parents[1]
FIXTURES = RAIZ / 'tests/fixtures'


@pytest.mark.parametrize('nome,rotulo', [
    ('sintetica_integra.xlsx', 'Lote'),
    ('sintetica_integra.xls', 'Número'),
    ('sintetica_corrompida.xls', 'Número'),
    ('sintetica_adulterada.xls', 'Número')])
def test_parser_le_rotulo_do_arquivo(nome, rotulo):
    parser = ExcelParser()
    try:
        parser.parse(FIXTURES / nome)
    except ErroLeituraExcel:
        if nome != 'sintetica_adulterada.xls': raise
    assert getattr(parser, 'rotulo_numero', None) == rotulo


def test_pipeline_e_parser_reutilizado_nao_vazam_rotulo():
    parser = ExcelParser()
    parser.parse(FIXTURES / 'sintetica_integra.xls')
    assert getattr(parser, 'rotulo_numero', None) == 'Número'
    parser.parse(FIXTURES / 'sintetica_integra.xlsx')
    assert getattr(parser, 'rotulo_numero', None) == 'Lote'
    assert getattr(processar_excel(FIXTURES / 'sintetica_integra.xls'), 'rotulo_numero', None) == 'Número'


@pytest.mark.parametrize('cabecalho', ['', '  Número  ', 'nÚmERo'])
def test_cabecalho_vazio_e_texto_exato(tmp_path, cabecalho):
    import openpyxl
    p = tmp_path / 'cabecalho_sintetico.xlsx'
    wb = openpyxl.Workbook()
    wb.active.append(['Data', cabecalho])
    wb.save(p); wb.close()
    parser = ExcelParser(); parser.parse(p)
    assert getattr(parser, 'rotulo_numero', None) == (cabecalho.strip() or 'Lote')


def test_sem_cabecalho_reinicia_lote(tmp_path):
    import openpyxl
    p = tmp_path / 'sem_cabecalho_sintetico.xlsx'
    wb = openpyxl.Workbook(); wb.active.append(['sem dados']); wb.save(p); wb.close()
    parser = ExcelParser(); parser.parse(FIXTURES / 'sintetica_integra.xls'); parser.parse(p)
    assert getattr(parser, 'rotulo_numero', None) == 'Lote'


def test_titulo_xlsx_valor_sem_api_nova():
    res = processar_excel(FIXTURES / 'sintetica_integra.xlsx').resultados[0]
    det = next(d for d in a.balde_detalhes(res.balde) if '500,00' in d['titulo'])
    assert det['titulo'] == 'Lote 3014 · 03/02/2025 · Crédito R\\$ 500,00 · Estorno sem a operação correspondente'
    assert '…' not in det['titulo'] and det['historico']


@pytest.mark.parametrize('numero,titulo', [
    ('3013','Número 3013 · 15/01/2025 · Débito R\\$ 3.000,00 · Sem NF correspondente'),
    ('3014','Número 3014 · 03/02/2025 · Crédito R\\$ 500,00 · Estorno sem a operação correspondente'),
    ('3016','Número 3016 · 01/04/2025 · Débito R\\$ 400,00 · Histórico não reconhecido')])
def test_rotulo_xls_titulo_e_identidade(numero, titulo):
    proc = processar_excel(FIXTURES / 'sintetica_integra.xls')
    assert getattr(proc, 'rotulo_numero', None) == 'Número'
    res = proc.resultados[0]
    detalhes = a.balde_detalhes(res.balde, rotulo_numero=proc.rotulo_numero)
    item = next(i for i in res.balde if i.lancamentos[0].numero == numero)
    det = next(d for d in detalhes if d['item_id'] == item.identificador)
    assert det['titulo'] == titulo
    assert det['historico'] == item.lancamentos[0].historico
    assert det['localizacao'].startswith(f'Número {numero} · ')


@pytest.mark.parametrize('banco,rotulo,esperado', [
    (False,'Número','Número'), (False,'Lote','Lote'), (False,'','Lote')])
def test_localizacao_rotulo_numero_vazio(banco, rotulo, esperado):
    res = processar_excel(FIXTURES / 'sintetica_integra.xls').resultados[0]
    l = next(l for i in res.balde for l in i.lancamentos if l.numero == '3014')
    assert a.localizacao_lancamento(l, 'Crédito', banco=banco, rotulo=rotulo) == f'{esperado} 3014 · Crédito · Contrapartida (Cta.C.Part.): 38'
    vazio = replace(l, numero='')
    assert a.localizacao_lancamento(vazio, 'Crédito', banco=banco, rotulo=rotulo) == f'{esperado} — · Crédito · Contrapartida (Cta.C.Part.): 38'


@pytest.mark.parametrize('n,origem,esperado', [
    (0,Origem.ALOCACAO_MANUAL,'NOVA: linha criada à mão nesta sessão'),
    (1,Origem.COMPRA_PERIODO,'NF NOVA: 1 quitação alocada à mão'),
    (2,Origem.COMPRA_PERIODO,'NF NOVA: 2 quitações alocadas à mão')])
def test_notas_manual_com_crase(n, origem, esperado):
    item = ComposicaoItem('NOVA', origem, Decimal('0'), Decimal('0'), Decimal('0'), Decimal('0'), quitacoes_manuais=n)
    assert a.notas_composicao_texto([item]) == [esperado]


def test_estorno_e_erro_do_core_sem_termo_interno():
    res = processar_excel(FIXTURES / 'sintetica_integra.xls').resultados[0]
    item = next(i for i in res.balde if i.lancamentos[0].numero == '3014')
    assert a.motivo_humano(item) == 'Histórico reconhecido como estorno, mas a operação que ele reverte não está nos itens sem destino — ou já foi conciliada com uma NF, ou não está nesta razão.'
    assert alocacao._MENSAGENS['ITEM_INEXISTENTE'] == 'este item não está mais entre os itens sem destino — o arquivo ou o motor mudou; a alocação ficou sem efeito'


def test_varredura_de_literais_de_tela():
    # Mesma extracao AST da regra de tela: nao conta comentario/docstring.
    from tests.test_regra_de_tela import _literais
    violacoes = []
    for p in (RAIZ / 'app').glob('*.py'):
        for _, s in _literais(p):
            # "balde" sozinho e chave programatica; frases sao texto de tela.
            if (re.search(r'\bbalde\b', s, re.I) and ' ' in s) or ' a mão' in s:
                violacoes.append((p.name, s))
    assert violacoes == []
    assert not any(re.search(r'\bbalde\b', s) for s in alocacao._MENSAGENS.values())


@pytest.mark.parametrize('nome,rotulo', [('sintetica_integra.xls','Número'), ('sintetica_integra.xlsx','Lote')])
def test_apptest_upload_titulo_caption_placeholder_e_chave(monkeypatch, nome, rotulo):
    at = AppTest.from_file(str(RAIZ / 'app/main.py'), default_timeout=60).run()
    at.file_uploader[0].set_value((nome, (FIXTURES/nome).read_bytes(), 'application/vnd.ms-excel')).run()
    assert not at.exception
    cx = [e for e in at.expander if 'cx::' in e.proto.id]
    res = processar_excel(FIXTURES/nome).resultados[0]
    assert len(cx) == len(res.balde) == 3
    assert {e.proto.id.split('cx::',1)[1] for e in cx} == {f'aloc::{nome}::7777::{i.identificador}' for i in res.balde}
    assert all(e.label.startswith(rotulo+' ') and '…' not in e.label for e in cx)
    assert f'{rotulo} 3014 · 03/02/2025 · Crédito R\\$ 500,00 · Estorno sem a operação correspondente' in [e.label for e in cx]
    assert any(c.value == f'{rotulo} 3014 · Crédito · Contrapartida (Cta.C.Part.): 38' for c in at.caption)
    assert all('(número do lançamento no Domínio)' not in c.value for c in at.caption)
    # o texto do campo diz "lote", a palavra do nome que o
    # motor da a linha criada em branco -- igual nos dois cabecalhos.
    item = next(i for i in res.balde if i.lancamentos[0].numero == '3014')
    campo = next(s for s in at.selectbox if s.key and s.key.endswith(item.identificador))
    campo.set_value(next(o for o in campo.options if 'nova' in o.lower())).run()
    assert not at.exception
    entrada = next(t for t in at.text_input if t.label == 'Nome da linha')
    assert entrada.proto.placeholder == 'opcional — em branco usa o lote'
    next(b for b in at.button if b.key == 'aplicar::' + entrada.key.split('rotulo::', 1)[1]).click().run()
    assert not at.exception
    assert [x for d in at.dataframe for x in d.value.astype(str).values.flatten() if 'manual' in x] == ['(manual: lote 3014)']
    assert 'ℹ️ (manual: lote 3014): linha criada à mão nesta sessão' in [c.value for c in at.caption]


# os testes do upload entram antes destes: mesma fixture local/sintetica, opt-in e finally.
from tests.test_texto_upload_e_dica import navegador_texto, pagina_texto


@pytest.mark.parametrize('nome,rotulo', [('sintetica_integra.xls','Número'), ('sintetica_integra.xlsx','Lote')])
def test_chromium_titulo_localizacao_motivo_e_aplicar(pagina_texto, nome, rotulo, tmp_path):
    import json
    from playwright.sync_api import expect
    page, tema = pagina_texto
    page.set_input_files('input[type="file"]', str(FIXTURES/nome))
    estorno = page.locator('details').filter(
        has=page.locator('summary').filter(has_text='Estorno sem a operação correspondente')).last
    expect(estorno).to_be_visible()
    esperado = f'{rotulo} 3014 · 03/02/2025 · Crédito R$ 500,00 · Estorno sem a operação correspondente'
    expect(estorno.locator('summary')).to_contain_text(esperado)
    medidas = page.locator('details summary').evaluate_all('''es => es.filter(e =>
        /Sem NF correspondente|Estorno sem a operação correspondente|Histórico não reconhecido/.test(e.innerText)).map(e => {
          const p=e.querySelector('p')||e; const r=p.getBoundingClientRect(); const s=getComputedStyle(p);
          return {titulo:p.innerText,largura:r.width,altura:r.height,overflow:s.textOverflow,
            scroll:p.scrollWidth,client:p.clientWidth,linhas:r.height/parseFloat(s.lineHeight)}; })''')
    (tmp_path/f'cabecalho_{tema}_{nome}.json').write_text(json.dumps(medidas,ensure_ascii=False,indent=2))
    print('TITULOS', tema, nome, json.dumps(medidas,ensure_ascii=False))
    assert len(medidas) == 3
    assert all('…' not in m['titulo'] and m['scroll'] <= m['client'] for m in medidas)
    assert all(m['titulo'].startswith(rotulo+' ') for m in medidas)
    estorno.locator('summary').click(); expect(estorno).to_have_attribute('open','')
    expect(estorno).to_contain_text(f'{rotulo} 3014 · Crédito · Contrapartida (Cta.C.Part.): 38')
    expect(estorno).to_contain_text('não está nos itens sem destino')
    assert '(número do lançamento no Domínio)' not in estorno.inner_text()
    page.screenshot(path=str(tmp_path/f'cabecalho_{tema}_{nome}.png'))
    # O debito 400,00 continua alocavel; nota e historico dentro intactos.
    debito = page.locator('details').filter(
        has=page.locator('summary').filter(has_text='Histórico não reconhecido')).last
    debito.locator('summary').click(); expect(debito).to_have_attribute('open','')
    expect(debito).to_contain_text('TARIFA ADMINISTRATIVA SEM PADRAO CONHECIDO')
    seletor = debito.locator('[data-testid="stSelectbox"] input')
    seletor.click(); page.get_by_role('option',name='NF 3333',exact=False).first.click()
    debito.get_by_role('button',name='Aplicar',exact=True).click()
    expect(page.get_by_text('NF 3333: 1 quitação alocada à mão',exact=False)).to_be_visible()
    # A frase exige uma declaracao fora; estorno fica disponivel.
    estorno = page.locator('details').filter(
        has=page.locator('summary').filter(has_text='Estorno sem a operação correspondente')).last
    estorno.locator('[data-testid="stSelectbox"] input').click()
    page.get_by_role('option').filter(has_text='não pertence').first.click()
    campo = estorno.locator('[data-testid="stTextInput"] input')
    campo.fill('justificativa sintetica do cabecalho')
    estorno.get_by_role('button',name='Aplicar',exact=True).click()
    expect(page.get_by_text('a diferença com o Domínio é essa.',exact=False)).to_be_visible()


def test_lote_de_uploads_preserva_rotulo_de_cada_arquivo(monkeypatch):
    at = AppTest.from_file(str(RAIZ/'app/main.py'), default_timeout=60).run()
    nomes = ['sintetica_integra.xls', 'sintetica_integra.xlsx']
    at.file_uploader[0].set_value([(n,(FIXTURES/n).read_bytes(),'application/vnd.ms-excel') for n in nomes]).run()
    assert not at.exception
    for n, rotulo in zip(nomes, ['Número','Lote']):
        cx = [e for e in at.expander if f'cx::aloc::{n}::' in e.proto.id]
        assert len(cx)==3 and all(e.label.startswith(rotulo+' ') for e in cx)


# --- segunda bateria do cabecalho ---------------------------------------------
_TITULO_3014 = '3014 · 03/02/2025 · Crédito R\\$ 500,00 · Estorno sem a operação correspondente'
_CINZA_3014 = '3014 · Crédito · Contrapartida (Cta.C.Part.): 38'
WJ = '\u2060'   # o caractere sem largura que desfaz o padrao de link automatico
_CAB_HOSTIL = ('http://exemplo.invalid Www.exemplo.invalid fulano@exemplo.invalid '
               '<https://exemplo.invalid> <ftp://exemplo.invalid> *x* [y](http://exemplo.invalid) R$')
_CAB_HOSTIL_NA_TELA = (f'http:{WJ}//exemplo.invalid Www{WJ}.exemplo.invalid fulano@{WJ}exemplo.invalid '
                       f'<{WJ}https:{WJ}//exemplo.invalid> <{WJ}ftp://exemplo.invalid> '
                       f'\\*x\\* \\[y\\](http:{WJ}//exemplo.invalid) R\\$')


# A3b (revisao do delta): o gatilho escrito como ENTIDADE HTML. As 9 formas medidas na
# revisao + outras da mesma classe (nomeada, decimal com zeros, hexadecimal em maiuscula).
_ENTIDADES = [
    'http&colon;//exemplo.invalid', 'http&#58;//exemplo.invalid', 'http:&#47;&#47;exemplo.invalid',
    '&#104;ttp://exemplo.invalid', 'www&period;exemplo.invalid', '&#119;ww.exemplo.invalid',
    'fulano&commat;exemplo.invalid', 'fulano&#64;exemplo.invalid',
    'http&#x3a;//exemplo.invalid www&#x2e;exemplo.invalid x&#x40;exemplo.invalid',
    'http&colon;&sol;&sol;exemplo.invalid', 'http&#0058;//exemplo.invalid', 'http&#X3A;//exemplo.invalid',
    '&#x77;ww&#46;exemplo.invalid', 'fulano&#x0040;exemplo.invalid', '&#60;http&#58;//exemplo.invalid&#62;']
_CAB_ENTIDADES = ' '.join(_ENTIDADES)
_MOTIVOS = 'Sem NF correspondente|Estorno sem a operação correspondente|Histórico não reconhecido'


def _item_3014():
    res = processar_excel(FIXTURES / 'sintetica_integra.xlsx').resultados[0]
    return next(i for i in res.balde if i.lancamentos[0].numero == '3014')


def _xlsx_com_cabecalho(tmp_path, cabecalho):
    """A sintetica_integra.xlsx com outro texto na celula do cabecalho da coluna do numero."""
    import openpyxl
    wb = openpyxl.load_workbook(FIXTURES / 'sintetica_integra.xlsx')
    next(r for r in wb.active.iter_rows() if r[0].value == 'Data')[1].value = cabecalho
    p = tmp_path / 'cabecalho_sintetico.xlsx'
    wb.save(p); wb.close()
    return p


@pytest.mark.parametrize('cabecalhos,esperado', [(['Lote', 'Número'], 'Lote'), (['Número', 'Lote'], 'Número')])
def test_cabecalho_do_primeiro_bloco_vale_pro_arquivo(tmp_path, cabecalhos, esperado):
    import openpyxl
    p = tmp_path / 'dois_blocos_sintetico.xlsx'
    wb = openpyxl.Workbook()
    for c in cabecalhos: wb.active.append(['Data', c])
    wb.save(p); wb.close()
    parser = ExcelParser(); parser.parse(p)
    assert parser.rotulo_numero == esperado


@pytest.mark.parametrize('rotulo', ['Lote', 'Número'])
def test_titulo_com_numero_vazio_sai_travessao(rotulo):
    item = _item_3014()
    vazio = replace(item, lancamentos=[replace(item.lancamentos[0], numero='')])
    det = a.balde_detalhes([vazio], rotulo_numero=rotulo)[0]
    assert det['titulo'] == f'{rotulo} — · 03/02/2025 · Crédito R\\$ 500,00 · Estorno sem a operação correspondente'
    assert det['localizacao'] == f'{rotulo} — · Crédito · Contrapartida (Cta.C.Part.): 38'


def test_rotulo_com_espacos_nas_pontas_sai_sem_eles():
    det = a.balde_detalhes([_item_3014()], rotulo_numero='  Número  ')[0]
    assert det['titulo'] == f'Número {_TITULO_3014}' and det['localizacao'] == f'Número {_CINZA_3014}'


def test_rotulo_do_arquivo_sai_escapado_no_titulo_e_na_linha_cinza():
    det = a.balde_detalhes([_item_3014()], rotulo_numero='Nº *do* [doc] _x_ `y` ~z~ ! R$')[0]
    escapado = 'Nº \\*do\\* \\[doc\\] \\_x\\_ \\`y\\` \\~z\\~ \\! R\\$'
    assert det['titulo'] == f'{escapado} {_TITULO_3014}'
    assert det['localizacao'] == f'{escapado} {_CINZA_3014}'


def test_rotulo_do_arquivo_nao_forma_link_automatico():
    det = a.balde_detalhes([_item_3014()], rotulo_numero=_CAB_HOSTIL)[0]
    assert det['titulo'] == f'{_CAB_HOSTIL_NA_TELA} {_TITULO_3014}'
    assert det['localizacao'] == f'{_CAB_HOSTIL_NA_TELA} {_CINZA_3014}'
    # cabecalho comum nao ganha caractere nenhum
    for rotulo in ('Lote', 'Número', 'Nº doc.', 'Núm: 1'):
        assert a.balde_detalhes([_item_3014()], rotulo_numero=rotulo)[0]['titulo'] == f'{rotulo} {_TITULO_3014}'


@pytest.mark.parametrize('entrada', _ENTIDADES + ['N&M', 'P&amp;L', '&'])
def test_e_comercial_do_rotulo_nao_deixa_entidade_ser_decodificada(entrada):
    # depois de todo '&' do cabecalho do arquivo vai o caractere sem largura: '&colon;' deixa de ser entidade
    na_tela = entrada.replace('&', '&' + WJ)
    det = a.balde_detalhes([_item_3014()], rotulo_numero=entrada)[0]
    assert det['titulo'] == f'{na_tela} {_TITULO_3014}'
    assert det['localizacao'] == f'{na_tela} {_CINZA_3014}'


@pytest.mark.parametrize('cabecalho,na_tela', [
    (_CAB_HOSTIL, _CAB_HOSTIL_NA_TELA), (_CAB_ENTIDADES, _CAB_ENTIDADES.replace('&', '&' + WJ))])
def test_apptest_cabecalho_hostil_do_arquivo_chega_neutralizado(monkeypatch, tmp_path, cabecalho, na_tela):
    p = _xlsx_com_cabecalho(tmp_path, cabecalho)
    assert processar_excel(p).rotulo_numero == cabecalho
    at = AppTest.from_file(str(RAIZ / 'app/main.py'), default_timeout=60).run()
    at.file_uploader[0].set_value((p.name, p.read_bytes(), 'application/vnd.ms-excel')).run()
    assert not at.exception
    cx = [e for e in at.expander if 'cx::' in e.proto.id]
    assert len(cx) == 3 and all(e.label.startswith(na_tela + ' ') for e in cx)
    assert f'{na_tela} {_TITULO_3014}' in [e.label for e in cx]
    assert f'{na_tela} {_CINZA_3014}' in [c.value for c in at.caption]


@pytest.mark.parametrize('cabecalho', [_CAB_HOSTIL, _CAB_ENTIDADES, 'N&M'], ids=['gatilhos', 'entidades', 'e_comercial'])
def test_chromium_cabecalho_hostil_sem_link_no_titulo_nem_na_linha_cinza(pagina_texto, tmp_path, cabecalho):
    page, tema = pagina_texto
    page.set_input_files('input[type="file"]', str(_xlsx_com_cabecalho(tmp_path, cabecalho)))
    caixas = page.locator('details').filter(has=page.locator('summary').filter(has_text=re.compile(_MOTIVOS)))
    caixas.nth(2).wait_for()
    for i in range(3):
        caixas.nth(i).locator('summary').click()
    caixas.nth(2).locator('[data-testid="stCaptionContainer"]').first.wait_for()
    medidas = caixas.evaluate_all("""es => es.map(d => {
        const s = d.querySelector('summary'), c = d.querySelector('[data-testid="stCaptionContainer"]');
        return {links_titulo: s.querySelectorAll('a').length, links_cinza: c.querySelectorAll('a').length,
                enfase: s.querySelectorAll('em,strong').length + c.querySelectorAll('em,strong').length,
                titulo: s.innerText, cinza: c.innerText}; })""")
    print('LINKS', tema, [(m['links_titulo'], m['links_cinza'], m['enfase']) for m in medidas])
    assert [(m['links_titulo'], m['links_cinza'], m['enfase']) for m in medidas] == [(0, 0, 0)] * 3
    # o que se le e o cabecalho do arquivo, letra por letra
    assert all(cabecalho in m['titulo'].replace(WJ, '') and cabecalho in m['cinza'].replace(WJ, '') for m in medidas)


# --- PAGINA CORTADA: o nome do arquivo no rotulo do seletor "Contas na tela" ---
_ARQUIVOS_HOSTIS = ['www.exemplo.invalid.xlsx', 'fulano@exemplo.invalid &amp; www&period;exemplo.invalid *b* _i_ [c].xlsx']
_ROTULOS_DO_SELETOR = [f'Contas na tela — www{WJ}.exemplo.invalid.xlsx',
                       f'Contas na tela — fulano@{WJ}exemplo.invalid &{WJ}amp; www&{WJ}period;exemplo.invalid '
                       '\\*b\\* \\_i\\_ \\[c\\].xlsx']


def _lote_de_arquivos_hostis():
    from tests.test_app_pagina_cortada import xlsx_de_n_contas
    return [(nome, xlsx_de_n_contas(6, 8001 + 100 * i)) for i, nome in enumerate(_ARQUIVOS_HOSTIS)]


def test_nome_do_arquivo_no_rotulo_do_seletor_chega_neutralizado(monkeypatch):
    at = AppTest.from_file(str(RAIZ / 'app/main.py'), default_timeout=120).run()
    at.file_uploader[0].set_value([(n, d, 'application/vnd.ms-excel') for n, d in _lote_de_arquivos_hostis()]).run()
    assert not at.exception
    assert sorted(s.label for s in at.selectbox if s.label.startswith('Contas na tela')) == sorted(_ROTULOS_DO_SELETOR)


def test_chromium_nome_do_arquivo_no_rotulo_do_seletor_sem_link_nem_formatacao(pagina_texto):
    from playwright.sync_api import expect
    page, tema = pagina_texto
    page.set_input_files('input[type="file"]', [
        {'name': n, 'mimeType': 'application/vnd.ms-excel', 'buffer': d} for n, d in _lote_de_arquivos_hostis()])
    seletores = page.locator('[data-testid="stSelectbox"]').filter(has_text='Contas na tela')
    expect(seletores).to_have_count(2)
    medidas = seletores.evaluate_all("""es => es.map(e => { const r = e.querySelector('label');
        return {links: r.querySelectorAll('a').length, enfase: r.querySelectorAll('em,strong').length, texto: r.innerText}; })""")
    print('SELETOR', tema, [(m['links'], m['enfase']) for m in medidas])
    assert [(m['links'], m['enfase']) for m in medidas] == [(0, 0), (0, 0)]
    # o que se le e o nome do arquivo, letra por letra
    assert sorted(m['texto'].replace(WJ, '') for m in medidas) == sorted(f'Contas na tela — {n}' for n in _ARQUIVOS_HOSTIS)
