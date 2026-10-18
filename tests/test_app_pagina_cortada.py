"""PAGINA CORTADA: no maximo K = 10 contas desenhadas por vez, triagem e planilha COMPLETAS e o fato
"Contas 1 a K de N deste arquivo" (seletor "Contas na tela"). Arquivos sinteticos de N contas gerados AQUI (N copias da
conta 7777 da sintetica_integra, codigos 8001...), nada de planilha real."""
import io
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from parser.excel import ExcelParser

RAIZ = Path(__file__).parent.parent
MAIN = RAIZ / "app" / "main.py"
INTEGRA = RAIZ / "tests" / "fixtures" / "sintetica_integra.xls"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ROTULO = "Contas na tela"
K = 10
ZERO = Decimal("0.00")
CAB = ["Data", "Número", "Histórico", None, None, None, None, "Cta.C.Part.", "Débito", "Crédito", None,
       "Saldo-Exercício"]


def xlsx_de_n_contas(n: int, primeiro: int = 8001) -> bytes:
    """N copias da conta 7777 da sintetica_integra (17 lancamentos, 3 itens sem destino cada), codigos primeiro..."""
    base = ExcelParser().parse(str(INTEGRA))[0]
    lancs = [(l.data, int(l.numero), l.historico, int(l.cta_part), l.debito, l.credito) for l in base.lancamentos]
    wb = openpyxl.Workbook()
    ws = wb.active
    for linha in (["Empresa:", None, "COMERCIAL EXEMPLO FICTICIA LTDA"], ["C.N.P.J.:", None, "00.000.000/0000-00"],
                  ["Período:", None, "01/01/2024 - 31/12/2025"], ["RAZÃO"], [], CAB):
        ws.append(linha)
    for i in range(n):
        ws.append(["Conta:", primeiro + i, "2.01.01.001.00", None, None, f"FORNECEDOR FICTICIO {i + 1:04d} LTDA"])
        ws.append([None, None, "SALDO ANTERIOR"] + [None] * 8 + [float(base.saldo_anterior)])
        saldo, td, tc = base.saldo_anterior, ZERO, ZERO
        for d, numero, hist, cta, deb, cre in lancs:
            saldo, td, tc = saldo + deb - cre, td + deb, tc + cre
            ws.append([datetime(d.year, d.month, d.day), numero, hist, None, None, None, None, cta,
                       float(deb) if deb != ZERO else None, float(cre) if cre != ZERO else None, None, float(saldo)])
        ws.append([])
        ws.append([None, None, None, None, "Total da conta:", None, None, None, float(td), float(tc)])
        ws.append([])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tela(*arquivos):
    at = AppTest.from_file(str(MAIN), default_timeout=120).run()
    at.file_uploader[0].set_value([(nome, dados, MIME) for nome, dados in arquivos]).run()
    return at


def _codigos(at):
    return [re.search(r"Conta (\d+)", s.value).group(1) for s in at.subheader if "Conta " in s.value]


def _seletor(at, rotulo=ROTULO):
    achados = [s for s in at.selectbox if s.label == rotulo]
    return achados[0] if achados else None


def _opcoes(total, janela=K):
    return [f"Contas {i + 1} a {min(i + janela, total)} de {total} deste arquivo" for i in range(0, total, janela)]


def test_75_contas_desenham_so_10_e_20_tabelas():
    at = _tela(("leve_75.xlsx", xlsx_de_n_contas(75)))
    assert len(at.exception) == 0
    assert _codigos(at) == [str(8001 + i) for i in range(10)]
    assert len(at.dataframe) == 20


def test_o_fato_e_o_seletor_com_as_8_janelas():
    at = _tela(("leve_75.xlsx", xlsx_de_n_contas(75)))
    seletor = _seletor(at)
    assert seletor is not None and seletor.value == "Contas 1 a 10 de 75 deste arquivo"
    assert list(seletor.options) == _opcoes(75) and seletor.options[-1] == "Contas 71 a 75 de 75 deste arquivo"


def test_a_triagem_continua_com_as_75_contas():
    at = _tela(("leve_75.xlsx", xlsx_de_n_contas(75)))
    triagem = next(m.value for m in at.markdown if "<table" in m.value and "Veredito" in m.value)
    assert triagem.count("<tr") - 1 == 75                      # 75 linhas + o cabecalho


def test_a_segunda_janela_mostra_8011_a_8020():
    at = _tela(("leve_75.xlsx", xlsx_de_n_contas(75)))
    _seletor(at).select("Contas 11 a 20 de 75 deste arquivo").run()
    assert len(at.exception) == 0
    assert _codigos(at) == [str(8011 + i) for i in range(10)]
    assert _seletor(at).value == "Contas 11 a 20 de 75 deste arquivo"


def test_a_ultima_janela_tem_5_contas():
    at = _tela(("leve_75.xlsx", xlsx_de_n_contas(75)))
    _seletor(at).select("Contas 71 a 75 de 75 deste arquivo").run()
    assert _codigos(at) == [str(8071 + i) for i in range(5)]


def test_a_planilha_leva_as_75_contas_com_a_1a_janela_na_tela():
    import tempfile

    from app.main import _downloads_do_lote, _EntradaLote
    from app.pipeline import processar_excel
    dados = xlsx_de_n_contas(75)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        f.write(dados)
    resultados = processar_excel(Path(f.name)).resultados
    xlsx = _downloads_do_lote([_EntradaLote(nome="leve_75.xlsx", resultados=resultados, dados=dados)])[0][2]
    aba = openpyxl.load_workbook(io.BytesIO(xlsx))["Conferência"]
    blocos = [c.value for linha in aba.iter_rows(max_col=1) for c in linha
              if isinstance(c.value, str) and re.match(r"Conta 80\d\d", c.value)]
    assert len(blocos) == 75


@pytest.mark.parametrize("n,tem_seletor", [(10, False), (11, True)])
def test_bordas_10_e_11(n, tem_seletor):
    at = _tela(("b.xlsx", xlsx_de_n_contas(n)))
    assert (_seletor(at) is not None) is tem_seletor
    assert len(_codigos(at)) == 10
    if tem_seletor:
        assert list(_seletor(at).options) == ["Contas 1 a 10 de 11 deste arquivo", "Contas 11 a 11 de 11 deste arquivo"]


def test_lote_de_2_arquivos_de_40_contas_divide_a_janela():
    at = _tela(("a.xlsx", xlsx_de_n_contas(40)), ("b.xlsx", xlsx_de_n_contas(40)))
    assert len(at.exception) == 0
    assert len(at.subheader) == 10                                # 5 de cada
    assert {s.label: s.value for s in at.selectbox if s.label.startswith(ROTULO)} == {
        f"{ROTULO} — a.xlsx": "Contas 1 a 5 de 40 deste arquivo", f"{ROTULO} — b.xlsx": "Contas 1 a 5 de 40 deste arquivo"}


def test_lote_de_10_arquivos_de_1_conta_nao_ganha_seletor():
    at = _tela(*[(f"u{i}.xlsx", xlsx_de_n_contas(1, 8001 + i)) for i in range(10)])
    assert len(at.exception) == 0 and _seletor(at) is None
    assert not any(s.label.startswith(ROTULO) for s in at.selectbox)
    assert len(at.subheader) == 10


def test_decisao_aplicada_na_janela_1_continua_la_depois_de_trocar_e_voltar():
    at = _tela(("leve_75.xlsx", xlsx_de_n_contas(75)))
    primeiro = next(s for s in at.selectbox if s.label == "Alocar em")
    primeiro.select_index(0).run()
    next(b for b in at.button if b.label == "Aplicar").click().run()
    assert any("Decisões desta sessão (1)" in m.value for m in at.markdown)
    _seletor(at).select("Contas 11 a 20 de 75 deste arquivo").run()
    assert not any("Decisões desta sessão (1)" in m.value for m in at.markdown)      # outra janela: sem a decisao
    _seletor(at).select("Contas 1 a 10 de 75 deste arquivo").run()
    assert any("Decisões desta sessão (1)" in m.value for m in at.markdown)
    assert len(at.exception) == 0


def test_o_texto_e_a_funcao_unica_e_passa_na_regra_de_tela():
    from app.apresentacao import contas_na_tela_texto
    from tests.test_regra_de_tela import _violacoes
    assert contas_na_tela_texto(1, 10, 75) == "Contas 1 a 10 de 75 deste arquivo"
    assert _violacoes([contas_na_tela_texto(1, 10, 75), contas_na_tela_texto(71, 75, 75), ROTULO,
                       f"{ROTULO} — arquivo.xlsx"]) == []


def test_reenvio_com_menos_contas_com_a_5a_janela_escolhida_nao_quebra():
    at = _tela(("leve.xlsx", xlsx_de_n_contas(75)))
    _seletor(at).select("Contas 41 a 50 de 75 deste arquivo").run()
    at.file_uploader[0].set_value([("leve.xlsx", xlsx_de_n_contas(12), MIME)]).run()
    assert len(at.exception) == 0
    assert _seletor(at).value == "Contas 1 a 10 de 12 deste arquivo"


def test_acrescentar_um_2o_arquivo_com_a_5a_janela_escolhida_volta_as_duas_a_1a():
    at = _tela(("a.xlsx", xlsx_de_n_contas(75)))
    _seletor(at).select("Contas 41 a 50 de 75 deste arquivo").run()
    at.file_uploader[0].set_value([("a.xlsx", xlsx_de_n_contas(75), MIME), ("b.xlsx", xlsx_de_n_contas(30), MIME)]).run()
    assert len(at.exception) == 0
    assert {s.label: s.value for s in at.selectbox if s.label.startswith(ROTULO)} == {
        f"{ROTULO} — a.xlsx": "Contas 1 a 5 de 75 deste arquivo", f"{ROTULO} — b.xlsx": "Contas 1 a 5 de 30 deste arquivo"}
