"""O registro do servico (stderr -> logs/conciliador_*.log) nao leva nome de arquivo nem nada da razao.

So fato neutro: texto fixo, contagem, numero de linha, TIPO da excecao, lugar do codigo e o hash curto (8 hex do sha256 dos
BYTES do arquivo). Tudo sintetico; a MARCA_* e' o que so existe no nome ou no conteudo do arquivo do teste.
A captura passa pelo handler que o servico instala (como tests/test_app_smoke.py faz) e, onde cabe, tambem pelo caplog (qualquer
outro handler): o registro nao pode carregar a mensagem da excecao por nenhum caminho. Avisos da biblioteca e o que o xlrd
imprime sao provados em PROCESSO FILHO (o pytest captura os avisos do processo dele e esconderia o defeito)."""
import ast
import hashlib
import io
import logging
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from parser.excel import ErroLeituraExcel, ExcelParser

RAIZ = Path(__file__).parent.parent
FIXTURES = RAIZ / "tests" / "fixtures"
CORROMPIDA = FIXTURES / "sintetica_corrompida.xls"
INTEGRA = FIXTURES / "sintetica_integra.xls"
DATA_HORA = 24                      # "2026-10-11 23:26:51,545 " = 24 caracteres


def h8(dados: bytes) -> str:
    return hashlib.sha256(dados).hexdigest()[:8]


@pytest.fixture
def servico(monkeypatch):
    """O que o handler do servico escreveria no stderr (o mesmo handler que app/main.py instala)."""
    import app.main  # noqa: F401  (instala os handlers de "app" e "parser")
    saida = io.StringIO()
    for nome in ("app", "parser"):
        (handler,) = logging.getLogger(nome).handlers
        monkeypatch.setattr(handler, "stream", saida)
    return saida


def _linhas(servico):
    return [linha[DATA_HORA:] for linha in servico.getvalue().splitlines()]


class _Upload:
    def __init__(self, nome, dados):
        self.name, self._dados = nome, dados

    def getbuffer(self):
        return self._dados


# === parser/excel.py: as 4 linhas de registro =======================================================

def test_reparo_leva_o_hash_e_nao_o_nome(servico):
    from app.pipeline import processar_excel
    processar_excel(CORROMPIDA, nome="MARCA_NOME.xls")
    assert _linhas(servico) == [f"INFO parser.excel: arquivo {h8(CORROMPIDA.read_bytes())}: reparado na leitura "
                                f"(BOUNDSHEET do exportador do Dominio)"]


def test_arquivo_que_nao_abre_leva_o_hash_e_so_o_tipo_do_leitor(servico, tmp_path):
    lixo = tmp_path / "tmpab12cd.xls"
    conteudo = b"MARCA_BYTES isto nao e um arquivo excel"
    lixo.write_bytes(conteudo)
    with pytest.raises(ErroLeituraExcel):
        ExcelParser().parse(lixo, nome="MARCA_NOME.xls")
    assert _linhas(servico) == [f"WARNING parser.excel: arquivo {h8(conteudo)}: nao abriu (leitor: XLRDError)"]


def test_xlsx_que_nao_abre_leva_o_hash_e_o_tipo(servico, tmp_path):
    quebrado = tmp_path / "tmpab12cd.xlsx"
    conteudo = b"PK\x03\x04 MARCA_BYTES isto nao e um zip de verdade"
    quebrado.write_bytes(conteudo)
    with pytest.raises(ErroLeituraExcel):
        ExcelParser().parse(quebrado, nome="MARCA_NOME.xlsx")
    assert _linhas(servico) == [f"WARNING parser.excel: arquivo {h8(conteudo)}: nao abriu como .xlsx (leitor: BadZipFile)"]


def test_linha_ilegivel_leva_o_hash_a_linha_e_o_tipo_sem_o_repr(servico, tmp_path):
    from tests.test_reparo_xls import _xls_com_celula_de_debito
    arquivo = _xls_com_celula_de_debito(tmp_path / "data_nan.xls", 10.0, data=float("nan"))
    with pytest.raises(ErroLeituraExcel):
        ExcelParser().parse(arquivo, nome="MARCA_NOME.xls")
    assert _linhas(servico) == [f"WARNING parser.excel: arquivo {h8(arquivo.read_bytes())}: linha 10 nao reconhecida "
                                f"(leitor: ValueError)"]


# === app/main.py: erro nao tratado = hash + TIPO + pilha, nunca a mensagem ==========================

def _erro_no_motor(monkeypatch):
    from app import main

    def _explode(*a, **k):
        raise RuntimeError("MARCA_MSG fornecedor 7341,59")
    monkeypatch.setattr(main, "processar_excel", _explode)
    monkeypatch.setattr(main.st, "session_state", {})
    return main


def test_erro_inesperado_processando_leva_hash_tipo_e_lugar_sem_nome_nem_mensagem(servico, monkeypatch, caplog):
    main = _erro_no_motor(monkeypatch)
    dados = b"qualquer coisa MARCA_BYTES"
    with caplog.at_level(logging.ERROR, logger="app.main"):
        (entrada,) = main._processar_lote([_Upload("MARCA_NOME.xls", dados)])
    assert entrada.erro == ("Erro inesperado processando 'MARCA_NOME.xls'; nenhum resultado foi gerado pra este arquivo.")
    texto = servico.getvalue()
    linhas = _linhas(servico)
    assert linhas[0] == f"ERROR app.main: erro inesperado processando o arquivo {h8(dados)}"
    assert "RuntimeError" in texto and "app/main.py" in texto
    for proibido in ("MARCA_MSG", "MARCA_NOME", "MARCA_BYTES", "7341", "Traceback"):
        assert proibido not in texto and proibido not in caplog.text, proibido


def test_planilha_de_um_arquivo_que_falha_leva_o_hash_sem_nome_nem_mensagem(servico, monkeypatch, caplog):
    from app import main
    from app.pipeline import processar_excel

    def _explode(*a, **k):
        raise RuntimeError("MARCA_MSG valor 9,99")
    monkeypatch.setattr(main, "gerar_excel", _explode)
    monkeypatch.setattr(main.st, "session_state", {})
    dados = INTEGRA.read_bytes()
    entrada = main._EntradaLote(nome="MARCA_NOME.xls", resultados=processar_excel(INTEGRA).resultados, dados=dados)
    falhas = []
    with caplog.at_level(logging.ERROR, logger="app.main"):
        assert main._downloads_do_lote([entrada], falhas) == [] and falhas == [entrada]
    assert _linhas(servico)[0] == f"ERROR app.main: erro inesperado gerando a planilha do arquivo {h8(dados)}"
    for proibido in ("MARCA_MSG", "MARCA_NOME", "9,99", "Traceback"):
        assert proibido not in servico.getvalue() and proibido not in caplog.text, proibido
    assert "RuntimeError" in servico.getvalue()


def test_clique_que_falha_leva_o_nome_da_funcao_o_tipo_e_nao_a_mensagem(servico, monkeypatch, caplog):
    from app import main
    monkeypatch.setattr(main.st, "session_state", {})

    @main._clique_protegido(lambda *a: "lugar")
    def aplicar_decisao():
        raise RuntimeError("MARCA_MSG historico MARCA_HIST")
    with caplog.at_level(logging.ERROR, logger="app.main"):
        aplicar_decisao()
    assert _linhas(servico)[0] == "ERROR app.main: erro inesperado na acao (aplicar_decisao)"
    for proibido in ("MARCA_MSG", "MARCA_HIST", "Traceback"):
        assert proibido not in servico.getvalue() and proibido not in caplog.text, proibido
    assert "RuntimeError" in servico.getvalue()


def test_erro_ao_montar_a_tela_leva_o_tipo_e_nao_a_mensagem(servico, monkeypatch, caplog):
    from streamlit.testing.v1 import AppTest

    import app.apresentacao as ap

    def _explode(*a, **k):
        raise RuntimeError("MARCA_MSG nome FORN_MARCA")
    monkeypatch.setattr(ap, "triagem_linhas", _explode)
    with caplog.at_level(logging.ERROR):
        at = AppTest.from_file(str(RAIZ / "app" / "main.py"), default_timeout=60).run()
        at.file_uploader[0].set_value([("MARCA_NOME.xls", INTEGRA.read_bytes(), "application/vnd.ms-excel")]).run()
    texto = servico.getvalue()
    assert "ERROR app.main: erro inesperado montando a tela" in _linhas(servico)
    for proibido in ("MARCA_MSG", "FORN_MARCA", "MARCA_NOME", "Traceback"):
        assert proibido not in texto and proibido not in caplog.text, proibido
    assert "RuntimeError" in texto


def test_nenhuma_chamada_de_registro_leva_traceback_nem_exc_info():
    """Trava de CLASSE: nenhum ponto novo de app/ ou parser/ volta a gravar a excecao inteira (que traz a mensagem)."""
    ruins = []
    for pasta in ("app", "parser"):
        for arquivo in sorted((RAIZ / pasta).glob("*.py")):
            for no in ast.walk(ast.parse(arquivo.read_text(encoding="utf-8"))):
                if isinstance(no, ast.Call) and isinstance(no.func, ast.Attribute):
                    if no.func.attr == "exception" or any(k.arg == "exc_info" for k in no.keywords):
                        ruins.append(f"{arquivo.name}:{no.lineno}")
    assert ruins == []


# === a hash: o mesmo valor no parser e no app =================================================

def test_o_hash_do_parser_e_o_do_app_sao_o_mesmo_pro_mesmo_arquivo(servico, monkeypatch):
    from app import main
    dados = CORROMPIDA.read_bytes()
    monkeypatch.setattr(main.st, "session_state", {})
    main._processar_lote([_Upload("MARCA_NOME.xls", dados)])                    # o parser grava a linha do reparo
    hash_do_parser = re.search(r"arquivo (\w{8}):", servico.getvalue()).group(1)
    servico.truncate(0)
    servico.seek(0)
    monkeypatch.setattr(main, "processar_excel", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("x")))
    main._processar_lote([_Upload("MARCA_NOME.xls", dados)])                    # o app grava a linha do erro
    hash_do_app = re.search(r"arquivo (\w{8})\n", servico.getvalue()).group(1)
    assert hash_do_parser == hash_do_app == h8(dados)


# === o que as bibliotecas escrevem sozinhas (processo filho) ========================================

def _filho(codigo: str, arquivo: Path):
    return subprocess.run([sys.executable, "-c", codigo, str(arquivo)], capture_output=True, text=True, cwd=RAIZ)


_CABECA = "import sys; sys.path.insert(0, '.'); from parser.excel import ExcelParser, grade_bruta; p = sys.argv[1]; "


@pytest.mark.parametrize("passo", [
    "ExcelParser().parse(p, nome='MARCA_NOME.xls')",
    "grade_bruta(open(p, 'rb').read())",
], ids=["parse", "grade_bruta"])
def test_o_xlrd_nao_escreve_no_stdout_do_servico(tmp_path, passo):
    arquivo = tmp_path / "tam.xls"
    arquivo.write_bytes(INTEGRA.read_bytes() + b"\x00" * 100)         # tamanho fora do multiplo de 512
    r = _filho(_CABECA + passo, arquivo)
    assert r.returncode == 0 and r.stdout == ""


@pytest.fixture
def xlsx_com_data_fora_do_limite(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 10):
        ws.cell(row=i, column=3, value="x")
    ws["A10"] = 987654321012                                         # a MARCA: o valor da celula
    ws["A10"].number_format = "dd/mm/yyyy"
    arquivo = tmp_path / "data.xlsx"
    wb.save(arquivo)
    return arquivo


@pytest.mark.parametrize("passo", [
    "ExcelParser().parse(p, nome='MARCA_NOME.xlsx')",
    "grade_bruta(open(p, 'rb').read())",
], ids=["parse", "grade_bruta"])
def test_aviso_do_openpyxl_nao_leva_o_conteudo_do_arquivo_pra_saida(xlsx_com_data_fora_do_limite, passo):
    r = _filho(_CABECA + passo, xlsx_com_data_fora_do_limite)
    assert r.returncode == 0
    assert "987654321012" not in r.stdout + r.stderr and "Cell A10" not in r.stdout + r.stderr
    assert (r.stdout + r.stderr).strip() == ""


def test_o_registro_leva_texto_fixo_categoria_contagem_e_hash(servico, xlsx_com_data_fora_do_limite):
    ExcelParser().parse(xlsx_com_data_fora_do_limite, nome="MARCA_NOME.xlsx")
    assert _linhas(servico) == [f"INFO parser.excel: arquivo {h8(xlsx_com_data_fora_do_limite.read_bytes())}: "
                                f"1 aviso do leitor nao gravado (UserWarning=1)"]


def test_avisos_calados_nao_deixam_o_filtro_de_avisos_alterado():
    import warnings
    antes = list(warnings.filters)
    from parser.registro import avisos_calados
    with avisos_calados(logging.getLogger("parser.excel"), lambda: "00000000"):
        warnings.warn("x", UserWarning)
    assert list(warnings.filters) == antes


# Trava de CLASSE (nota aditiva A3): todo argumento que vai pra uma linha de registro (`_log.x(...)`, `_LOG.x(...)` e
# `registrar_excecao(...)`) esta numa lista fechada de fatos neutros. Argumento novo (um nome de arquivo, uma mensagem de
# excecao, um valor da razao) derruba o teste ate alguem decidir de proposito e acrescentar aqui.
_ARGUMENTOS_NEUTROS = {                  # por ARQUIVO: "erro" e "n" valem num lugar so; o mesmo nome em outro arquivo e diferente
    "excel.py": {"hash_do_arquivo(source)", "type(erro).__name__", "type(e).__name__", "r + 1"},
    "main.py": {"hash_curto(uploaded.getbuffer())", "hash_curto(entrada.dados)", "funcao.__name__", "n",
                "'conta' if n == 1 else 'contas'"},
}
_NOMES_DE_REGISTRO = {"_log", "_LOG"}


def test_toda_chamada_de_registro_so_leva_argumento_da_lista_de_fatos_neutros():
    fora = []
    for pasta in ("app", "parser"):
        for arquivo in sorted((RAIZ / pasta).glob("*.py")):
            if arquivo.name == "registro.py":           # o proprio ajudante (monta a linha a partir de fatos ja neutros)
                continue
            for no in ast.walk(ast.parse(arquivo.read_text(encoding="utf-8"))):
                if not isinstance(no, ast.Call):
                    continue
                f = no.func
                if isinstance(f, ast.Attribute) and isinstance(f.value, ast.Name) and f.value.id in _NOMES_DE_REGISTRO:
                    mensagem, *args = no.args
                elif isinstance(f, ast.Name) and f.id == "registrar_excecao":
                    _, mensagem, *args = no.args
                else:
                    continue
                if not isinstance(mensagem, ast.Constant) or not isinstance(mensagem.value, str):
                    fora.append(f"{arquivo.name}:{no.lineno} mensagem nao literal")
                fora += [f"{arquivo.name}:{no.lineno} {ast.unparse(a)}" for a in args
                         if ast.unparse(a) not in _ARGUMENTOS_NEUTROS.get(arquivo.name, set())]
    assert fora == []


def test_uso_normal_nao_escreve_nada_no_registro(servico, monkeypatch):
    """CONTROLE (trava): enviar arquivo integro e baixar a planilha = 0 linha."""
    from app.main import _downloads_do_lote, _EntradaLote
    from app.pipeline import processar_excel
    entrada = _EntradaLote(nome="MARCA_NOME.xls", resultados=processar_excel(INTEGRA).resultados, dados=INTEGRA.read_bytes())
    assert len(_downloads_do_lote([entrada])) == 1
    assert servico.getvalue() == ""


def test_a_trava_nao_aceita_por_nome_o_que_leva_dado():
    """A mesma palavra vale em UM arquivo so: `erro` em parser/excel.py (o texto do leitor, com bytes) nao passa."""
    assert "erro" not in _ARGUMENTOS_NEUTROS["excel.py"] and "n" not in _ARGUMENTOS_NEUTROS["excel.py"]
    assert "erro" not in _ARGUMENTOS_NEUTROS["main.py"] and "causa" not in _ARGUMENTOS_NEUTROS["main.py"]
    assert not set.intersection(*_ARGUMENTOS_NEUTROS.values())


# === a trava dos avisos ====================================================================

@pytest.fixture(autouse=True)
def _avisos_ja_registrados_zerados(monkeypatch):
    from parser import registro
    monkeypatch.setattr(registro, "_AVISOS_JA_REGISTRADOS", set(), raising=False)


def _xlsx_com_datas_fora_do_limite(caminho, quantas):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 10):
        ws.cell(row=i, column=3, value="x")
    for i in range(quantas):
        celula = ws.cell(row=10 + i, column=1, value=987654321012 + i)
        celula.number_format = "dd/mm/yyyy"
    wb.save(caminho)
    return caminho


def test_chamada_aninhada_na_mesma_thread_nao_trava_o_servico():
    """Em PROCESSO FILHO com prazo: se a trava voltar a ser nao reentrante o filho trava e o pytest nao fica pendurado."""
    codigo = ("import sys, logging; sys.path.insert(0, '.'); from parser.registro import avisos_calados; "
              "log = logging.getLogger('parser.excel')\n"
              "with avisos_calados(log, lambda: '00000000'):\n"
              "    with avisos_calados(log, lambda: '00000000'):\n"
              "        print('dentro')\n")
    try:
        r = subprocess.run([sys.executable, "-c", codigo], capture_output=True, text=True, cwd=RAIZ, timeout=15)
    except subprocess.TimeoutExpired:
        pytest.fail("chamada aninhada travou (trava nao reentrante)")
    assert r.returncode == 0 and r.stdout.strip() == "dentro"


def test_duas_leituras_ao_mesmo_tempo_cada_uma_conta_os_seus_avisos(servico, tmp_path, capfd):
    import threading
    import warnings
    a = _xlsx_com_datas_fora_do_limite(tmp_path / "a.xlsx", 1)
    b = _xlsx_com_datas_fora_do_limite(tmp_path / "b.xlsx", 2)
    filtros = list(warnings.filters)
    largada = threading.Barrier(2)

    def ler(caminho):
        largada.wait(5)
        ExcelParser().parse(caminho, nome="MARCA_NOME.xlsx")
    fios = [threading.Thread(target=ler, args=(c,)) for c in (a, b)]
    for f in fios:
        f.start()
    for f in fios:
        f.join(30)
    assert not any(f.is_alive() for f in fios)
    assert sorted(_linhas(servico)) == sorted([
        f"INFO parser.excel: arquivo {h8(a.read_bytes())}: 1 aviso do leitor nao gravado (UserWarning=1)",
        f"INFO parser.excel: arquivo {h8(b.read_bytes())}: 2 avisos do leitor nao gravados (UserWarning=2)"])
    saida = capfd.readouterr()
    assert "98765432101" not in saida.out + saida.err + servico.getvalue() and "Cell A1" not in saida.out + saida.err
    assert list(warnings.filters) == filtros


def test_o_xls_nao_espera_a_trava_dos_avisos_e_o_xlsx_espera(tmp_path):
    import threading

    from parser import registro
    from parser.excel import grade_bruta
    xlsx = _xlsx_com_datas_fora_do_limite(tmp_path / "w.xlsx", 1)
    resultados = {}

    def ler(nome, dados):
        grade_bruta(dados)
        resultados[nome] = True
    with registro._TRAVA_DOS_AVISOS:                    # outra leitura de .xlsx esta com a trava
        fio_xls = threading.Thread(target=ler, args=("xls", INTEGRA.read_bytes()), daemon=True)
        fio_xlsx = threading.Thread(target=ler, args=("xlsx", xlsx.read_bytes()), daemon=True)
        fio_xls.start()
        fio_xlsx.start()
        fio_xls.join(5)
        assert resultados == {"xls": True}              # o .xls passou sem esperar
        assert fio_xlsx.is_alive()                      # o .xlsx ficou em fila (o estado de avisos e do processo todo)
    fio_xlsx.join(5)
    assert resultados == {"xls": True, "xlsx": True}


def test_a_linha_de_aviso_sai_uma_vez_por_arquivo_nao_por_leitura(servico, xlsx_com_data_fora_do_limite):
    from parser.excel import grade_bruta
    ExcelParser().parse(xlsx_com_data_fora_do_limite, nome="MARCA_NOME.xlsx")
    grade_bruta(xlsx_com_data_fora_do_limite.read_bytes())                     # a copia fiel do mesmo arquivo
    ExcelParser().parse(xlsx_com_data_fora_do_limite, nome="MARCA_NOME.xlsx")  # a reexecucao da tela
    assert len(_linhas(servico)) == 1


# === pilha e xlrd da rota de reparo ===============================================

def _pilha(excecao):
    from parser.registro import pilha_sem_mensagem
    return pilha_sem_mensagem(excecao)


def test_a_mensagem_da_causa_e_do_contexto_nao_sai_na_pilha():
    def com_causa():
        try:
            raise KeyError("MARCA_CAUSA fornecedor")
        except KeyError as k:
            raise ValueError("MARCA_MSG") from k

    def com_contexto():
        try:
            raise OSError("MARCA_CONTEXTO valor 7341,59")
        except OSError:
            raise RuntimeError("MARCA_MSG2")

    def suprimida():
        try:
            raise KeyError("MARCA_SUPRIMIDA")
        except KeyError:
            raise ValueError("MARCA_MSG3") from None
    for funcao, tipo_da_origem in ((com_causa, "KeyError"), (com_contexto, "OSError")):
        with pytest.raises(Exception) as e:
            funcao()
        texto = _pilha(e.value)
        assert f"causa: {tipo_da_origem} em " in texto and "MARCA" not in texto
    with pytest.raises(ValueError) as e:
        suprimida()
    assert "causa:" not in _pilha(e.value) and "MARCA" not in _pilha(e.value)


def test_nota_anexada_e_grupo_de_excecoes_nao_levam_mensagem():
    erro = ValueError("MARCA_MSG")
    erro.add_note("MARCA_NOTA")
    assert "MARCA" not in _pilha(erro)
    with pytest.raises(ExceptionGroup) as e:
        raise ExceptionGroup("MARCA_GRUPO", [ValueError("MARCA_SUB")])
    assert "MARCA" not in _pilha(e.value) and _pilha(e.value).startswith("ExceptionGroup em ")


def test_quadro_de_biblioteca_mostra_pasta_arquivo_linha_sem_o_caminho_da_instalacao(tmp_path):
    import traceback
    import zipfile
    with pytest.raises(Exception) as e:
        openpyxl.load_workbook(io.BytesIO(b"MARCA_BYTES isto nao e um zip"))
    texto = _pilha(e.value)
    # O quadro da biblioteca padrao sai do PROPRIO Python que roda o teste: no 3.13 o `zipfile` e uma pasta
    # (`zipfile/__init__.py`), no 3.11 e um arquivo (`python3.11/zipfile.py`). Garante que esse quadro faz parte da pilha
    # e reduz o caminho real a pasta/arquivo, sem usar o codigo de produto pra montar o esperado.
    arquivos_da_pilha = set()
    erro = e.value
    while erro is not None:
        arquivos_da_pilha |= {q.filename for q in traceback.extract_tb(erro.__traceback__)}
        erro = erro.__cause__ or erro.__context__
    assert zipfile.__file__ in arquivos_da_pilha
    pasta_e_arquivo = "/".join(Path(zipfile.__file__).parts[-2:])
    inicio_do_quadro = r"(?:em|<-) "                                   # o quadro comeca no nome: pasta/arquivo, nada antes
    for lugar in (pasta_e_arquivo, "openpyxl/reader/excel.py"):
        assert re.search(inicio_do_quadro + re.escape(lugar) + r":\d+ \(", texto), lugar
    assert zipfile.__file__ not in texto and str(Path(zipfile.__file__).parent.parent) not in texto
    for proibido in (sys.prefix, sys.base_prefix, str(Path.home()), "site-packages", "MARCA_BYTES"):
        assert proibido not in texto, proibido
    # modulo de fora do repo e da instalacao: so pasta/arquivo
    import importlib.util
    pasta = tmp_path / "pasta_de_fora"
    pasta.mkdir()
    (pasta / "modulo_de_fora.py").write_text("def explode():\n    raise RuntimeError('MARCA_X')\n", encoding="utf-8")
    spec = importlib.util.spec_from_file_location("modulo_de_fora", pasta / "modulo_de_fora.py")
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    with pytest.raises(RuntimeError) as e:
        modulo.explode()
    texto = _pilha(e.value)
    assert "pasta_de_fora/modulo_de_fora.py:2 (explode)" in texto and str(tmp_path) not in texto


@pytest.mark.parametrize("instalacao", [RAIZ / ".venv", Path("/opt/pasta_de_fora/venv")], ids=["venv_dentro", "venv_fora"])
def test_quadro_de_biblioteca_sai_sem_o_caminho_da_venv_com_ela_dentro_ou_fora_da_pasta(instalacao):
    # O teste acima so prova a venv em que a suite roda; aqui o quadro leva o caminho de uma biblioteca instalada numa
    # venv DENTRO e numa FORA da pasta do projeto, qualquer que seja a venv desta rodada (nenhum arquivo e criado).
    arquivo = instalacao.resolve() / "lib" / "python3.13" / "site-packages" / "pacote_ficticio" / "modulo.py"
    escopo = {}
    exec(compile("def explode():\n    raise RuntimeError('MARCA_X')\n", str(arquivo), "exec"), escopo)
    with pytest.raises(RuntimeError) as e:
        escopo["explode"]()
    texto = _pilha(e.value)
    assert re.search(r"(?:em|<-) pacote_ficticio/modulo\.py:2 \(explode\)", texto), texto
    for proibido in (str(instalacao), ".venv", "site-packages", "python3.13", "MARCA_X"):
        assert proibido not in texto, proibido


def test_toda_abertura_do_xlrd_inclusive_a_do_reparo_leva_destino_descartavel(monkeypatch):
    import xlrd

    from parser.excel import grade_bruta
    chamadas = []
    original = xlrd.open_workbook

    def espia(*args, **kwargs):
        chamadas.append(kwargs.get("logfile"))
        return original(*args, **kwargs)
    monkeypatch.setattr(xlrd, "open_workbook", espia)
    ExcelParser().parse(CORROMPIDA, nome="MARCA_NOME.xls")      # 1a abertura falha e a do REPARO abre
    grade_bruta(CORROMPIDA.read_bytes())                        # idem na copia fiel
    assert len(chamadas) == 4
    assert all(isinstance(c, io.StringIO) for c in chamadas)


def test_nenhuma_chamada_de_xlrd_open_workbook_do_parser_fica_sem_logfile():
    sem = []
    for arquivo in sorted((RAIZ / "parser").glob("*.py")):
        for no in ast.walk(ast.parse(arquivo.read_text(encoding="utf-8"))):
            if (isinstance(no, ast.Call) and isinstance(no.func, ast.Attribute) and no.func.attr == "open_workbook"
                    and isinstance(no.func.value, ast.Name) and no.func.value.id == "xlrd"
                    and not any(k.arg == "logfile" for k in no.keywords)):
                sem.append(f"{arquivo.name}:{no.lineno}")
    assert sem == []
