"""gerar_excel: 1 .xlsx por arquivo processado, DUAS abas.

Aba 1 "Razão (Domínio)" = copia FIEL do arquivo que o usuario subiu (reaberto
dos bytes ORIGINAIS, nao reconstruido do parseado) — a prova de fidelidade e o
que o parser IGNORA: rodape de licenciamento e a composicao manual do contador
nas ultimas linhas do golden.
Aba 2 "Conferência" = espelho da tela, com dinheiro em celula NUMERICA (o
conferente trabalha em cima: soma, filtra, marca) e data em celula de data.
"""
import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.pipeline import processar_excel
from relatorio.excel import ABA_CONFERENCIA, ABA_RAZAO, CLASSIFICACOES, COL_CONFERIDO, gerar_excel
from tests.helpers import GOLDEN, identidade_das_linhas

FIXTURES = Path(__file__).parent / "fixtures"
CORROMPIDA = FIXTURES / "sintetica_corrompida.xls"
INTEGRA_XLSX = FIXTURES / "sintetica_integra.xlsx"
GERADO_EM = datetime(2026, 7, 28, 14, 0)


def _excel(caminho: Path, nome: str | None = None) -> bytes:
    res = processar_excel(caminho).resultados
    return gerar_excel(res, caminho.read_bytes(), nome or caminho.name,
                       gerado_em=GERADO_EM)


def _abas(dados: bytes):
    return openpyxl.load_workbook(io.BytesIO(dados))


def _celulas(ws) -> list:
    return [c for linha in ws.iter_rows() for c in linha if c.value is not None]


def _textos(ws) -> list[str]:
    return [c.value for c in _celulas(ws) if isinstance(c.value, str)]


def _numeros(ws) -> list[Decimal]:
    return [Decimal(str(c.value)) for c in _celulas(ws)
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


def _celula_composicao(ws, nf: str, rotulo: str):
    """A CELULA da linha da NF `nf` na coluna de cabecalho `rotulo` da tabela de
    composicao — por NOME, nunca por indice: coluna nova no meio da tabela
    nao pode derrubar teste que nao fala de posicao (mesma licao do wrap por
    nome)."""
    cab = next(c.row for c in _celulas(ws) if c.value == "NF")
    linha = next(c.row for c in _celulas(ws)
                 if c.value == nf and c.column == 1 and c.row > cab)
    return ws.cell(row=linha, column=_coluna(ws, cab, rotulo))


def _linha_que_comeca_com(ws, valor) -> list:
    """A primeira linha cuja PRIMEIRA celula preenchida e `valor` (as tabelas
    da aba comecam na coluna A) -> lista de valores da linha."""
    for linha in ws.iter_rows(values_only=True):
        if linha and linha[0] == valor:
            return list(linha)
    raise AssertionError(f"nenhuma linha comeca com {valor!r}")


# --- estrutura: duas abas, na ordem -------------------------------------------

def test_gerar_excel_do_golden_tem_as_duas_abas_na_ordem():
    wb = _abas(_excel(GOLDEN))
    assert wb.sheetnames == [ABA_RAZAO, ABA_CONFERENCIA]


def test_arquivo_gerado_e_um_xlsx_de_verdade():
    dados = _excel(GOLDEN)
    assert dados.startswith(b"PK")                  # container ZIP do xlsx
    assert len(dados) > 1000


# --- aba 1: copia FIEL do original --------------------------------------------

def test_aba_razao_copia_o_que_o_parser_ignora_no_original():
    """Rodape de licenciamento e a composicao manual do contador (ultimas
    linhas do golden) NAO entram na list[Conta] — se estao na aba, ela veio do
    ARQUIVO, nao da reconstrucao."""
    ws = _abas(_excel(GOLDEN))[ABA_RAZAO]
    textos = _textos(ws)
    assert "Sistema licenciado para ESCRITORIO FICTICIO LTDA" in textos
    assert "composição" in textos
    assert "nf 618204" in textos
    assert "INDUSTRIA FICTICIA DE ALIMENTOS LTDA" in textos   # preambulo
    assert "Período:" in textos


def test_aba_razao_preserva_o_numero_cru_do_original_com_artefato_de_float():
    """A ultima linha da composicao manual do golden traz 1440.090000000001
    na celula. Copia FIEL = o numero que esta la, nao o arredondado."""
    ws = _abas(_excel(GOLDEN))[ABA_RAZAO]
    assert Decimal("1440.090000000001") in _numeros(ws)


def test_aba_razao_tem_o_mesmo_numero_de_linhas_do_original():
    import xlrd
    original = xlrd.open_workbook(str(GOLDEN)).sheet_by_index(0)
    ws = _abas(_excel(GOLDEN))[ABA_RAZAO]
    assert ws.max_row == original.nrows


def test_aba_razao_converte_data_declarada_do_biff_em_data():
    """Coluna Data do golden e' XL_CELL_DATE (serial + formato de data): vira
    celula de DATA, nunca o serial 45311 solto."""
    ws = _abas(_excel(GOLDEN))[ABA_RAZAO]
    datas = [c.value for c in _celulas(ws) if isinstance(c.value, datetime)]
    assert date(2024, 1, 20) in [d.date() for d in datas]   # 1o lancamento


def test_aba_razao_do_xlsx_original_tambem_sai_fiel():
    """Rota openpyxl (magic bytes PK): mesma copia fiel."""
    ws = _abas(_excel(INTEGRA_XLSX))[ABA_RAZAO]
    assert "FORNECEDOR FICTICIO LTDA" in " ".join(_textos(ws))


def test_aba_razao_do_xls_corrompido_usa_a_rota_de_reparo():
    """O .xls que o Dominio exporta corrompido (BOUNDSHEET) e' o caso de uso
    comum: a aba fiel passa pela MESMA rota de reparo do parser, senao o
    arquivo que a ferramenta consegue conferir nao consegue exportar."""
    ws = _abas(_excel(CORROMPIDA))[ABA_RAZAO]
    assert "FORNECEDOR FICTICIO LTDA" in " ".join(_textos(ws))


# --- aba 2: o golden ao centavo, em celula NUMERICA ----------------------------

def test_secao_de_decisoes_manuais_so_aparece_quando_ha_decisao():
    """O que o conferente decidiu vai pro cliente com evidencia. Planilha
    do caso comum (ninguem decidiu nada) nao ganha um "Nenhuma" a mais."""
    from core.modelo import AlocacaoManual, DestinoAlocacao

    sintetica = Path(__file__).parent / "fixtures" / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]

    sem = openpyxl.load_workbook(io.BytesIO(gerar_excel(
        [antes], sintetica.read_bytes(), "s.xls")))[ABA_CONFERENCIA]
    assert not any(t.startswith("Decisões manuais") for t in _textos(sem))

    decididos = processar_excel(sintetica, {antes.conta.codigo: [AlocacaoManual(
        item_id=antes.balde[0].identificador,
        destino=DestinoAlocacao.FORA_DA_COMPOSICAO,
        rotulo="lançamento e de outra conta")]}).resultados
    com = openpyxl.load_workbook(io.BytesIO(gerar_excel(
        decididos, sintetica.read_bytes(), "s.xls")))[ABA_CONFERENCIA]
    textos = _textos(com)
    assert "Decisões manuais (1)" in textos      # a contagem no titulo
    assert any("fora da composição" in t for t in textos)
    assert any("lançamento e de outra conta" in t for t in textos)
    # o historico integral acompanha a decisao (e a evidencia pro cliente)
    historico = decididos[0].decisoes[0].item.lancamentos[0].historico
    assert any(historico in t for t in textos)


def test_wrap_das_decisoes_segue_a_coluna_do_historico_e_nao_o_indice():
    """A coluna Origem entrou ANTES do Historico e reindexou a tabela. Com
    wrap por indice o Historico teria perdido a quebra de linha em silencio (a
    evidencia sai numa tira ilegivel) sem derrubar teste nenhum. Este teste e o
    que faltava — trava o wrap na coluna CERTA, seja qual for a posicao dela."""
    from core.modelo import AlocacaoManual, CriterioSugestao, DestinoAlocacao

    sintetica = Path(__file__).parent / "fixtures" / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]
    # com criterio_sugerido a coluna Origem EXISTE (coluna vazia some, e
    # e' justamente com ela presente que o risco de reindexacao aparece)
    decididos = processar_excel(sintetica, {antes.conta.codigo: [AlocacaoManual(
        item_id=antes.balde[0].identificador,
        destino=DestinoAlocacao.FORA_DA_COMPOSICAO,
        rotulo="lançamento e de outra conta",
        criterio_sugerido=CriterioSugestao.SERIE_X_VALOR_DECLARADO)]}).resultados
    ws = openpyxl.load_workbook(io.BytesIO(gerar_excel(
        decididos, sintetica.read_bytes(), "s.xls")))[ABA_CONFERENCIA]

    cab = next(c.row for c in _celulas(ws) if c.value == "Origem")
    historico = _coluna(ws, cab, "Histórico")
    origem = _coluna(ws, cab, "Origem")
    assert ws.cell(row=cab + 1, column=historico).alignment.wrap_text is True
    assert not ws.cell(row=cab + 1, column=origem).alignment.wrap_text


def test_wrap_por_nome_inexistente_estoura_em_vez_de_sumir():
    """Nome errado no wrap e erro de programacao: tem que quebrar no teste, nao
    virar uma coluna sem quebra de linha que ninguem percebe."""
    from relatorio.excel import _Folha

    folha = _Folha(openpyxl.Workbook().active)
    with pytest.raises(ValueError):
        folha.tabela(("A", "B"), [{"A": 1, "B": 2}], wrap=("Coluna Fantasma",))


def test_aba_conferencia_traz_a_composicao_do_golden_como_numero():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    abertos = [_celula_composicao(ws, nf, "EM ABERTO").value
               for nf in ("618204", "274591")]
    # numero, NUNCA texto: o conferente soma a coluna no proprio Excel
    assert all(isinstance(v, (int, float)) and not isinstance(v, bool)
               for v in abertos)
    assert [Decimal(str(v)) for v in abertos] == [Decimal("2820.00"), Decimal("1440.09")]
    assert Decimal(str(_celula_composicao(ws, "618204", "Valor compra").value)) == \
        Decimal("4700.00")


def test_aba_conferencia_traz_o_saldo_final_do_dominio_como_numero():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    linha = _linha_que_comeca_com(ws, "Saldo final (Domínio)")
    assert isinstance(linha[1], (int, float))
    assert Decimal(str(linha[1])) == Decimal("-4260.09")


def test_valores_de_dinheiro_saem_com_formato_brasileiro():
    """#,##0.00 e o codigo INVARIANTE do OOXML (a virgula e o ponto sao
    marcadores de posicao): o Excel pt-BR do escritorio renderiza 1.234,56."""
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    rotulo = next(c for c in _celulas(ws) if c.value == "Saldo final (Domínio)")
    assert ws.cell(row=rotulo.row, column=2).number_format == "#,##0.00"
    assert _celula_composicao(ws, "618204", "EM ABERTO").number_format == "#,##0.00"


def test_datas_da_composicao_saem_como_data_nao_texto():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    data_nf = _celula_composicao(ws, "618204", "Data NF").value
    assert isinstance(data_nf, datetime)
    assert data_nf.date() == date(2024, 12, 11)


def test_composicao_termina_com_coluna_conferido_vazia():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    cabecalho = [v for v in _linha_que_comeca_com(ws, "NF") if v is not None]
    assert cabecalho[-1] == "Conferido"
    assert _celula_composicao(ws, "618204", "Conferido").value is None   # o conferente e quem marca


def test_cabecalho_traz_arquivo_conta_periodo_e_saldos():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    texto = " ".join(_textos(ws))
    assert "golden_sintetico.xls" in texto
    assert "28/07/2026 14:00" in texto
    assert "Conta 5218 — FERRAMENTAS FICTICIAS COMERCIO SA" in texto
    assert "30/04/2025" in texto                      # fim do periodo
    assert Decimal(str(_linha_que_comeca_com(ws, "Saldo anterior")[1])) == Decimal("-429.26")
    assert "Fecha: a composição bate com o saldo final" in texto


def test_nenhuma_aba_sai_protegida():
    """O conferente trabalha EM CIMA do arquivo: nada travado."""
    wb = _abas(_excel(GOLDEN))
    for ws in wb.worksheets:
        assert ws.protection.sheet is False


# --- balde: aparece integral, com o historico ---------------------------------

def test_balde_sai_integral_na_aba_conferencia():
    res = processar_excel(CORROMPIDA).resultados[0]
    dados = gerar_excel([res], CORROMPIDA.read_bytes(), "sintetica_corrompida.xls",
                        gerado_em=GERADO_EM)
    ws = _abas(dados)[ABA_CONFERENCIA]
    texto = " ".join(_textos(ws))
    assert res.balde
    for item in res.balde:
        for l in item.lancamentos:
            assert l.historico in texto               # evidencia INTEGRAL
            assert l.numero in texto                  # lote


# --- multi-conta: blocos empilhados na mesma aba, ordem da tela ----------------

def test_multiconta_empilha_os_blocos_na_ordem_da_tela():
    from core.modelo import Conta, ResultadoConferencia
    res = [ResultadoConferencia(conta=Conta(codigo="1111", classificacao="2.01",
                                            nome="ALFA FICTICIA LTDA")),
           ResultadoConferencia(conta=Conta(codigo="2222", classificacao="2.01",
                                            nome="BETA FICTICIA LTDA"))]
    ws = _abas(gerar_excel(res, GOLDEN.read_bytes(), "duas_contas.xls"))[ABA_CONFERENCIA]
    texto = " ".join(_textos(ws))
    assert texto.index("ALFA FICTICIA LTDA") < texto.index("BETA FICTICIA LTDA")


# --- nome NFD do upload: normaliza NFC antes de renderizar ------------

def test_nome_de_arquivo_nfd_sai_normalizado_no_cabecalho():
    # upload do macOS chega DECOMPOSTO (a + til combinante U+0303); as formas
    # ficam em escape explicito — o literal identico esconderia a diferenca
    from core.modelo import Conta, ResultadoConferencia
    res = [ResultadoConferencia(conta=Conta(codigo="1", classificacao="x", nome="X"))]
    ws = _abas(gerar_excel(res, GOLDEN.read_bytes(),
                           "Raza\u0303o.xls"))[ABA_CONFERENCIA]
    assert "Raz\u00e3o.xls" in " ".join(_textos(ws))


# === FASE 2 — classificacao do balde (dropdown) ===============================

def _balde_do_corrompida():
    """(worksheet, linha do cabecalho do balde, qtd de linhas de dado)."""
    res = processar_excel(CORROMPIDA).resultados[0]
    dados = gerar_excel([res], CORROMPIDA.read_bytes(),
                        "sintetica_corrompida.xls", gerado_em=GERADO_EM)
    ws = _abas(dados)[ABA_CONFERENCIA]
    cabecalho = next(c.row for c in _celulas(ws) if c.value == "Lote")
    return ws, cabecalho, sum(len(i.lancamentos) for i in res.balde)


def _coluna(ws, cab, rotulo):
    """A coluna (1-based) de um rotulo do cabecalho do balde. Buscar pelo NOME
    em vez de fixar a letra: coluna nova no balde nao pode quebrar o teste do
    dropdown, so' o teste da POSICAO (abaixo), que e onde a regra mora."""
    return next(c for c in range(1, ws.max_column + 1)
                if ws.cell(row=cab, column=c).value == rotulo)


def test_balde_tem_colunas_de_classificacao_observacao_e_conferido():
    ws, cab, _ = _balde_do_corrompida()
    primeira = _coluna(ws, cab, "Classificação")
    rotulos = [ws.cell(row=cab, column=primeira + i).value for i in range(3)]
    assert rotulos == ["Classificação", "Observação", "Conferido"]


def test_sugestao_do_motor_fica_colada_a_esquerda_da_classificacao():
    """A proposta do motor e o dropdown que o conferente preenche ficam
    lado a lado — e onde ele trabalha na planilha."""
    ws, cab, _ = _balde_do_corrompida()
    esquerda = _coluna(ws, cab, "Classificação") - 1
    assert ws.cell(row=cab, column=esquerda).value == "Sugestão do motor"


def test_colunas_de_trabalho_do_balde_vem_vazias():
    """O sistema nao chuta a classificacao: quem preenche e o conferente."""
    ws, cab, n = _balde_do_corrompida()
    primeira = _coluna(ws, cab, "Classificação")
    for r in range(cab + 1, cab + 1 + n):
        assert [ws.cell(row=r, column=primeira + i).value
                for i in range(3)] == [None] * 3


def test_cada_linha_do_balde_tem_dropdown_de_classificacao():
    ws, cab, n = _balde_do_corrompida()
    validacoes = ws.data_validations.dataValidation
    assert validacoes
    letra = get_column_letter(_coluna(ws, cab, "Classificação"))
    for r in range(cab + 1, cab + 1 + n):
        assert any(f"{letra}{r}" in dv.sqref
                   for dv in validacoes), f"linha {r} sem dropdown"


def test_dropdown_traz_as_opcoes_combinadas():
    ws, _, _ = _balde_do_corrompida()
    dv = ws.data_validations.dataValidation[0]
    assert dv.type == "list"
    for opcao in CLASSIFICACOES:
        assert opcao in dv.formula1
    assert len(CLASSIFICACOES) == 6


def test_dropdown_e_sugestao_nao_camisa_de_forca():
    """Pode ficar vazio e aceita texto fora da lista (a lista aprende com o
    uso; recusar o que o conferente escreveu perderia informacao)."""
    ws, _, _ = _balde_do_corrompida()
    dv = ws.data_validations.dataValidation[0]
    assert dv.showErrorMessage is False
    assert dv.allowBlank is True


def test_instrucao_curta_acima_do_balde():
    ws, cab, _ = _balde_do_corrompida()
    acima = " ".join(str(c.value) for c in _celulas(ws) if c.row < cab)
    assert "Coluna Classificação" in acima          # descreve a coluna
    assert "melhoria do sistema" in acima


def test_composicao_nao_ganha_dropdown():
    """A validacao cobre SO o balde: marcar 'Conferido' na composicao e' texto
    livre, nao ha lista de classificacao la."""
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    linha_nf = next(c.row for c in _celulas(ws) if c.value == "618204")
    for dv in ws.data_validations.dataValidation:
        assert f"I{linha_nf}" not in dv.sqref


# === FASE 3 — materialidade, contagem e ordem por valor =======================

def test_balde_do_excel_sai_ordenado_por_valor_decrescente():
    """A TELA segue cronologica (#28 wont-fix); no Excel o que pesa vem no
    topo — la a coluna Data permite reordenar com um clique."""
    ws, cab, n = _balde_do_corrompida()
    valores = [ws.cell(row=r, column=4).value for r in range(cab + 1, cab + 1 + n)]
    # o credito de 250,00 (transferencia de saldos) saiu do balde
    # e virou linha da composicao (decisao de produto)
    assert valores == [3000, 500, 400]


def test_ordem_do_balde_e_a_mesma_em_duas_geracoes():
    """Mesmo arquivo = mesma planilha (requisito de produto)."""
    def _lotes():
        ws, cab, n = _balde_do_corrompida()
        return [ws.cell(row=r, column=1).value for r in range(cab + 1, cab + 1 + n)]
    assert _lotes() == _lotes()


def test_materialidade_e_contagem_saem_acima_do_balde():
    ws, cab, _ = _balde_do_corrompida()
    acima = " ".join(str(c.value) for c in _celulas(ws) if c.row < cab)
    # balde de 3 itens cabe nos "5 maiores" -- sem a frase
    assert "maiores itens respondem por" not in acima
    # o credito de 250,00 (transferencia de saldos) saiu do balde
    # e virou linha da composicao (decisao de produto)
    assert "3 itens: 1 Estorno sem par, 1 Não reconhecido, 1 Sem NF" in acima


def test_composicao_do_golden_nao_ganha_linha_de_materialidade():
    """Balde vazio: nada a resumir — a linha nao ocupa espaco."""
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    assert "maiores itens respondem" not in " ".join(_textos(ws))


# === o valor alocado a mao chega a planilha ============================

def test_alocacao_manual_chega_a_coluna_alocado_a_mao_da_planilha():
    """Regressao: o que o contador alocou na tela sai na
    planilha como NUMERO, na coluna propria — e "Total pago" fica so' com o que
    o motor casou. Mesma fonte da tela (composicao_celulas): um caminho so'."""
    from core.modelo import AlocacaoManual, DestinoAlocacao

    sintetica = FIXTURES / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]
    item = next(i for i in antes.balde if i.lancamentos[0].numero == "3013")   # 3.000,00
    res = processar_excel(sintetica, {antes.conta.codigo: [AlocacaoManual(
        item_id=item.identificador, destino=DestinoAlocacao.LINHA_ABERTA,
        alvo="1111")]}).resultados
    ws = _abas(gerar_excel(res, sintetica.read_bytes(), "s.xls"))[ABA_CONFERENCIA]
    alocado = _celula_composicao(ws, "1111", "Alocado à mão")
    assert Decimal(str(alocado.value)) == Decimal("3000.00")
    assert alocado.number_format == "#,##0.00"
    assert Decimal(str(_celula_composicao(ws, "1111", "Total pago").value)) == \
        Decimal("1200.00")
    assert Decimal(str(_celula_composicao(ws, "1111", "EM ABERTO").value)) == \
        Decimal("-1100.00")
    assert _celula_composicao(ws, "1111", "Última quitação").value.date() == \
        date(2025, 1, 15)


def _linhas_da_composicao(ws) -> list[dict]:
    """As linhas de DADO da tabela de composicao (a PRIMEIRA tabela com
    cabecalho "NF") como dicts cabecalho -> valor."""
    return _linhas_da_tabela(ws, next(c.row for c in _celulas(ws) if c.value == "NF"))


def _linhas_da_tabela(ws, cab: int) -> list[dict]:
    """As linhas de DADO da tabela cujo cabecalho esta na linha `cab` (a tabela
    acaba onde "EM ABERTO" fica vazio: as notas so' tem a coluna A. Nao
    "Valor compra", que fica vazio nas linhas que nao sao nota)."""
    colunas = {ws.cell(row=cab, column=c).value: c
               for c in range(1, ws.max_column + 1) if ws.cell(row=cab, column=c).value}
    linhas = []
    r = cab + 1
    while ws.cell(row=r, column=colunas["EM ABERTO"]).value is not None:
        linhas.append({nome: ws.cell(row=r, column=c).value for nome, c in colunas.items()})
        r += 1
    return linhas


def _excel_da_sintetica_com_alocacao():
    from core.modelo import AlocacaoManual, DestinoAlocacao

    sintetica = FIXTURES / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]
    item = next(i for i in antes.balde if i.lancamentos[0].numero == "3013")   # 3.000,00
    res = processar_excel(sintetica, {antes.conta.codigo: [AlocacaoManual(
        item_id=item.identificador, destino=DestinoAlocacao.LINHA_ABERTA,
        alvo="1111")]}).resultados
    return _abas(gerar_excel(res, sintetica.read_bytes(), "s.xls"))[ABA_CONFERENCIA]


def test_planilha_traz_retido_em_coluna_propria():
    sintetica = FIXTURES / "sintetica_integra.xls"
    ws = _abas(_excel(sintetica))[ABA_CONFERENCIA]
    retido = _celula_composicao(ws, "3333", "Retido")
    assert Decimal(str(retido.value)) == Decimal("150.00")
    assert retido.number_format == "#,##0.00"


def test_planilha_fecha_a_identidade_da_tabela_em_toda_linha():
    """Em CELULA numerica, com a familia inteira em jogo: retencao (3333),
    baixa no anterior (10.000,00 contra o saldo trazido) e o alocado a mao na
    1111 — que fica negativa e mesmo assim fecha."""
    linhas = _linhas_da_composicao(_excel_da_sintetica_com_alocacao())
    assert {l["NF"] for l in linhas} >= {"1111", "3333", "6666", "(saldo anterior)"}
    identidade_das_linhas(linhas, de=lambda v: Decimal(str(v)))


# === as NFs quitadas na planilha ================================

def _linha_do_texto(ws, inicio: str) -> int:
    return next(c.row for c in _celulas(ws)
                if isinstance(c.value, str) and c.value.startswith(inicio))


def _cabecalho_das_quitadas(ws) -> int:
    titulo = _linha_do_texto(ws, "NFs quitadas no período")
    return next(c.row for c in _celulas(ws) if c.value == "NF" and c.row > titulo)


def test_planilha_traz_as_quitadas_na_mesma_tabela_da_composicao():
    """UMA tabela (as quitadas nao ganham bloco proprio), com
    as quitadas em EM ABERTO zero, na ordenacao unica por idade. Um caminho so'
    entre tela e planilha."""
    ws = _excel_da_sintetica_com_alocacao()
    assert not any(isinstance(c.value, str) and c.value.startswith("NFs quitadas")
                   for c in _celulas(ws))
    linhas = _linhas_da_composicao(ws)
    # "(transferencia lote 3015)" entra pela idade (303 dias), entre
    # 5555 e 6666 — o credito de 250,00 virou linha da composicao
    assert [l["NF"] for l in linhas] == ["1111", "2222", "3333", "4444", "5555",
                                         "(transferência lote 3015)", "6666",
                                         "(saldo anterior)",
                                          "TOTAL"]          # o rodape
    abertos = {l["NF"]: Decimal(str(l["EM ABERTO"])) for l in linhas}
    assert [abertos[nf] for nf in ("2222", "4444", "5555")] == [Decimal("0.00")] * 3
    por_nf = {l["NF"]: l for l in linhas}
    assert Decimal(str(por_nf["4444"]["Total pago"])) == Decimal("650.00")  # multi-NF
    assert all(l["Conferido"] is None for l in linhas)
    identidade_das_linhas(linhas, de=lambda v: Decimal(str(v)))


# === a sugestao na planilha ========================================

def _excel_com_sugestao(monkeypatch, alocacoes=None):
    """A planilha da sintetica com uma regra de sugestao falsa ligada: o que se
    trava aqui e a PLANILHA, nao o criterio."""
    from core.modelo import CriterioSugestao, DestinoAlocacao, Sugestao

    def regra_falsa(item, ctx):
        abertas = [nf for nf in ctx.nfs
                   if nf.data_compra is not None and nf.saldo_aberto > 0]
        return None if not abertas else Sugestao(
            destino=DestinoAlocacao.LINHA_ABERTA, alvo=abertas[0].numero,
            criterio=CriterioSugestao.SERIE_X_VALOR_DECLARADO,
            evidencia={"parcelas": Decimal("2"),
                       "parcela_declarada": Decimal("541.37"),
                       "valor_compra": Decimal("1082.74"),
                       "pago": Decimal("554.07"),
                       "diferenca": Decimal("12.70")})

    monkeypatch.setattr("core.sugestao.REGRAS", (regra_falsa,))
    sintetica = Path(__file__).parent / "fixtures" / "sintetica_integra.xls"
    res = processar_excel(sintetica, alocacoes or {}).resultados
    dados = gerar_excel(res, sintetica.read_bytes(), "s.xls")
    return openpyxl.load_workbook(io.BytesIO(dados))[ABA_CONFERENCIA], res


def test_planilha_do_balde_traz_a_frase_da_sugestao(monkeypatch):
    """Quem confere le a proposta na linha do item, colada no dropdown que
    preenche — sem precisar abrir a tela."""
    ws, _ = _excel_com_sugestao(monkeypatch)
    cab = next(c.row for c in _celulas(ws) if c.value == "Lote")
    coluna = _coluna(ws, cab, "Sugestão do motor")
    frases = [ws.cell(row=cab + r, column=coluna).value for r in (1, 2)]
    assert any(f and "541,37" in f and "1.082,74" in f for f in frases)


def test_frase_da_sugestao_quebra_linha_na_planilha(monkeypatch):
    """A frase e uma sentenca: sem wrap ela sai numa tira que empurra o
    dropdown pra fora da tela."""
    ws, _ = _excel_com_sugestao(monkeypatch)
    cab = next(c.row for c in _celulas(ws) if c.value == "Lote")
    coluna = _coluna(ws, cab, "Sugestão do motor")
    assert ws.cell(row=cab + 1, column=coluna).alignment.wrap_text is True


def test_decisoes_declaram_a_sugestao_aceita(monkeypatch):
    """O rastro que chega ao cliente: esta linha existe porque o conferente
    aceitou uma proposta do motor."""
    from core.modelo import AlocacaoManual, CriterioSugestao, DestinoAlocacao

    ws, _ = _excel_com_sugestao(monkeypatch)          # so' pra achar o item
    sintetica = Path(__file__).parent / "fixtures" / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]
    aberta = next(i.nf_numero for i in antes.itens if i.saldo_aberto > 0)
    ws, _ = _excel_com_sugestao(monkeypatch, {antes.conta.codigo: [
        AlocacaoManual(item_id=antes.balde[0].identificador,
                       destino=DestinoAlocacao.LINHA_ABERTA, alvo=aberta,
                       criterio_sugerido=CriterioSugestao.SERIE_X_VALOR_DECLARADO)]})
    cab = next(c.row for c in _celulas(ws) if c.value == "Origem")
    assert ws.cell(row=cab + 1,
                   column=_coluna(ws, cab, "Origem")).value == (
        "sugestão do motor aceita")


def test_toda_coluna_do_balde_tem_largura_definida(monkeypatch):
    """A coluna nova nao pode ficar sem largura: no Excel ela sai no default
    estreito e a frase fica ilegivel."""
    from relatorio.excel import _LARGURAS

    ws, _ = _excel_com_sugestao(monkeypatch)
    cab = next(c.row for c in _celulas(ws) if c.value == "Lote")
    ultima = _coluna(ws, cab, COL_CONFERIDO)
    assert len(_LARGURAS) >= ultima


# --- a palavra "gap" nao existe pro usuario, tambem na planilha ---

def test_a_palavra_gap_nao_aparece_na_planilha():
    import re

    ws = _excel_da_sintetica_com_alocacao()
    achados = [c.value for c in _celulas(ws)
               if isinstance(c.value, str) and re.search(r"\bgaps?\b", c.value, re.I)]
    assert achados == []


# === planilha enxuta ===========================================

def _tabela_apos(ws, titulo: str, primeira_coluna: str) -> list[dict]:
    """As linhas da tabela cujo cabecalho comeca com `primeira_coluna`, logo
    apos o texto `titulo`."""
    linha_titulo = _linha_do_texto(ws, titulo)
    cab = next(c.row for c in _celulas(ws)
               if c.value == primeira_coluna and c.row > linha_titulo)
    return _linhas_da_tabela_ate_vazio(ws, cab)


def _linhas_da_tabela_ate_vazio(ws, cab: int) -> list[dict]:
    colunas = {ws.cell(row=cab, column=c).value: c
               for c in range(1, ws.max_column + 1) if ws.cell(row=cab, column=c).value}
    linhas, r = [], cab + 1
    while any(ws.cell(row=r, column=c).value is not None for c in colunas.values()):
        linhas.append({nome: ws.cell(row=r, column=c).value for nome, c in colunas.items()})
        r += 1
    return linhas


def test_cabecalho_da_tabela_de_erros_contabeis_tem_acento():
    """A coluna do nome da verificacao e' 'Verificação' (a tela e a
    planilha leem a mesma chave de divergencias_celulas)."""
    from core.modelo import Achado, Nivel
    res = processar_excel(GOLDEN).resultados
    res[0].achados.append(Achado(Nivel.ERRO, "self_check_saldo_final", "não fecha",
                                 esperado=Decimal("150.00"),
                                 encontrado=Decimal("100.00")))
    dados = gerar_excel(res, GOLDEN.read_bytes(), GOLDEN.name, gerado_em=GERADO_EM)
    ws = _abas(dados)[ABA_CONFERENCIA]
    cabecalho = next(c.row for c in _celulas(ws)
                     if c.value == "Mensagem"
                     and c.row > _linha_do_texto(ws, "Erros contábeis"))
    colunas = [c.value for c in ws[cabecalho] if c.value is not None]
    assert colunas == ["Verificação", "Mensagem", "Esperado", "Encontrado"]


def test_planilha_sem_alerta_nenhum_nao_ganha_a_secao():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    assert not any(t.startswith("Alertas (") for t in _textos(ws))


def test_decisoes_sem_justificativa_nem_origem_nao_levam_as_colunas():
    """Coluna cheia de '—' e' coluna que so' ocupa largura. So'
    aparece a que tem conteudo em ALGUMA linha."""
    from core.modelo import AlocacaoManual, DestinoAlocacao

    sintetica = FIXTURES / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]
    item = next(i for i in antes.balde if i.lancamentos[0].numero == "3013")
    res = processar_excel(sintetica, {antes.conta.codigo: [AlocacaoManual(
        item_id=item.identificador, destino=DestinoAlocacao.LINHA_ABERTA,
        alvo="1111")]}).resultados
    ws = _abas(gerar_excel(res, sintetica.read_bytes(), "s.xls"))[ABA_CONFERENCIA]
    linhas = _tabela_apos(ws, "Decisões manuais", "Lote")
    assert list(linhas[0]) == ["Lote", "Data", "Lado", "Valor", "Decisão",
                               "Histórico"]
    assert "3.000" in str(linhas[0]["Valor"]) or linhas[0]["Valor"] == 3000


def test_decisao_com_justificativa_mantem_a_coluna():
    """A recíproca: quem declarou item fora da composicao PRECISA da coluna --
    a justificativa e' o que vai pro cliente."""
    from core.modelo import AlocacaoManual, DestinoAlocacao

    sintetica = FIXTURES / "sintetica_integra.xls"
    antes = processar_excel(sintetica).resultados[0]
    item = next(i for i in antes.balde if i.lancamentos[0].numero == "3016")
    res = processar_excel(sintetica, {antes.conta.codigo: [AlocacaoManual(
        item_id=item.identificador, destino=DestinoAlocacao.FORA_DA_COMPOSICAO,
        rotulo="tarifa do banco, não e do fornecedor")]}).resultados
    ws = _abas(gerar_excel(res, sintetica.read_bytes(), "s.xls"))[ABA_CONFERENCIA]
    linhas = _tabela_apos(ws, "Decisões manuais", "Lote")
    assert "Justificativa" in linhas[0]
    assert linhas[0]["Justificativa"] == "tarifa do banco, não e do fornecedor"
    assert "Origem" not in linhas[0]          # nao houve sugestao aceita


_PROSA_CORTADA_DA_PLANILHA = (
    "O que o conferente decidiu na tela",
    "precisa de conferência humana",
)


def test_planilha_nao_carrega_prosa_entre_as_secoes():
    """O que se le uma vez vive no
    manual. Fica o que e' ASK (a instrucao da coluna Classificacao)."""
    ws = _excel_da_sintetica_com_alocacao()
    textos = " ".join(_textos(ws))
    assert [p for p in _PROSA_CORTADA_DA_PLANILHA if p in textos] == []
    assert any("Coluna Classificação" in t for t in _textos(ws))


def test_coluna_de_referencia_some_quando_nenhum_alerta_tem_uma():
    """Regra: coluna vazia em TODAS as linhas nao ocupa largura.
    Conta montada AQUI (cobertura de CI, sem depender de razao real): saldo
    anterior que ninguem detalhou -> alerta da CONTA, sem NF nem lote."""
    from datetime import date as _date

    from core.conferencia import conferir
    from core.modelo import Conta, Lancamento
    from parser.historico import ClassificadorHistorico

    lanc = Lancamento(data=_date(2026, 1, 10), numero="L1",
                      historico="Vr. compras conf. 555 FORNECEDOR FICTICIO LTDA",
                      cta_part="1.1.1", debito=Decimal("0.00"),
                      credito=Decimal("100.00"),
                      saldo_exercicio=Decimal("-1100.00"))
    conta = Conta(codigo="9999", classificacao="2.01.01", nome="FORNECEDOR FICTICIO",
                  data_fim_periodo=_date(2026, 3, 31),
                  saldo_anterior=Decimal("-1000.00"), lancamentos=[lanc],
                  total_debito_dominio=Decimal("0.00"),
                  total_credito_dominio=Decimal("100.00"))
    res = conferir(conta, ClassificadorHistorico())
    sintetica = FIXTURES / "sintetica_integra.xls"
    ws = _abas(gerar_excel([res], sintetica.read_bytes(), "s.xls"))[ABA_CONFERENCIA]

    linhas = _tabela_apos(ws, "Alertas", "Tipo")
    assert [l["Tipo"] for l in linhas] == ["Saldo anterior não detalhado"]
    assert "NF/Lote" not in linhas[0]


# === linha TOTAL e ponte do saldo na planilha =========================

def _linha_do_rotulo(ws, rotulo: str) -> int:
    return next(c.row for c in _celulas(ws) if c.value == rotulo and c.column == 1)


def test_composicao_do_golden_fecha_com_a_linha_total_em_negrito():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    linhas = _linhas_da_composicao(ws)
    assert [l["NF"] for l in linhas] == ["618204", "274591", "TOTAL"]
    total = linhas[-1]
    assert Decimal(str(total["Valor compra"])) == Decimal("7100.03")
    assert Decimal(str(total["Total pago"])) == Decimal("2839.94")
    assert Decimal(str(total["EM ABERTO"])) == Decimal("4260.09")
    assert total["Data NF"] is None and total["Idade (dias)"] is None
    identidade_das_linhas([total], de=lambda v: Decimal(str(v)))
    celula = _celula_composicao(ws, "TOTAL", "EM ABERTO")
    assert celula.font.bold
    assert celula.fill.fgColor.rgb not in (None, "00000000")
    assert celula.number_format == "#,##0.00"


def _ponte_da_planilha(ws) -> dict:
    """rotulo (coluna A) -> (valor da coluna E, marca da coluna F), da linha
    seguinte ao titulo "Ponte do saldo" ate a nota (se houver) ou a primeira
    linha vazia."""
    r = _linha_do_texto(ws, "Ponte do saldo") + 1
    ponte = {}
    while (ws.cell(row=r, column=1).value is not None
           and not str(ws.cell(row=r, column=1).value).startswith("Nota:")):
        ponte[ws.cell(row=r, column=1).value] = (ws.cell(row=r, column=5).value,
                                                 ws.cell(row=r, column=6).value)
        r += 1
    return ponte


def test_planilha_traz_a_ponte_do_golden_com_numero_e_marca():
    from tests.test_app_ponte import ROTULOS_PONTE

    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    ponte = _ponte_da_planilha(ws)
    assert list(ponte) == ROTULOS_PONTE            # mesma ordem da tela, zeros inclusos
    assert Decimal(str(ponte["(+) Compras do período"][0])) == Decimal("7100.03")
    assert Decimal(str(ponte["(-) Retenções"][0])) == Decimal("0.00")
    assert Decimal(str(ponte["(=) Saldo final apurado"][0])) == Decimal("4260.09")
    assert ponte["Saldo final Domínio"] == (4260.09, "OK")
    apurado = _linha_do_rotulo(ws, "(=) Saldo final apurado")
    assert ws.cell(row=apurado, column=1).font.bold
    assert ws.cell(row=apurado, column=5).number_format == "#,##0.00"
    assert not ws.cell(row=_linha_do_rotulo(ws, "(-) Retenções"), column=1).font.bold


def test_planilha_explica_o_total_que_nao_bate_com_a_ponte_so_quando_difere():
    golden = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    assert any("429,26 foi pago no período" in t for t in _textos(golden))
    sintetica = _abas(_excel(FIXTURES / "sintetica_integra.xls"))[ABA_CONFERENCIA]
    assert not any("foi pago no período" in t for t in _textos(sintetica))
    assert "Ponte do saldo" in _textos(sintetica)


# === Fechado por saldo zero na planilha ==========================

def _excel_que_zera(monkeypatch, tmp_path):
    from tests.test_app_corte import _xlsx_que_zera
    arquivo = tmp_path / "zera.xlsx"
    arquivo.write_bytes(_xlsx_que_zera(monkeypatch))
    return _abas(_excel(arquivo))[ABA_CONFERENCIA]


def test_planilha_traz_o_periodo_zerado_logo_apos_a_ponte(monkeypatch, tmp_path):
    ws = _excel_que_zera(monkeypatch, tmp_path)
    ponte = _linha_do_texto(ws, "Ponte do saldo")
    titulo = _linha_do_texto(ws, "Fechado por saldo zero")
    erros = _linha_do_texto(ws, "Erros contábeis")
    assert ponte < titulo < erros
    assert any("não divergência" in t for t in _textos(ws))

    nfs = _linhas_da_tabela_ate_vazio(
        ws, next(c.row for c in _celulas(ws) if c.value == "NF" and c.row > titulo))
    assert [l["NF"] for l in nfs] == ["5001", "5002"]
    assert [Decimal(str(l["EM ABERTO"])) for l in nfs] == [Decimal("700.00"),
                                                           Decimal("300.00")]
    assert all(l["Zerou em"].date() == date(2025, 1, 5) for l in nfs)

    sem_nf = _linhas_da_tabela_ate_vazio(
        ws, next(c.row for c in _celulas(ws) if c.value == "Lote" and c.row > titulo))
    assert [(l["Lote"], Decimal(str(l["Valor"]))) for l in sem_nf] == [
        ("4003", Decimal("1000.00"))]


def test_alerta_do_periodo_zerado_entra_na_tabela_de_alertas(monkeypatch, tmp_path):
    """Alerta, pra que planilha e triagem o enxerguem."""
    ws = _excel_que_zera(monkeypatch, tmp_path)
    tipos = [l["Tipo"] for l in _tabela_apos(ws, "Alertas", "Tipo")]
    assert "Período já zerado" in tipos


def test_planilha_sem_periodo_zerado_pendente_nao_traz_a_secao():
    ws = _abas(_excel(GOLDEN))[ABA_CONFERENCIA]
    assert "Fechado por saldo zero" not in _textos(ws)


# === a planilha fala como a tela ===================================

def test_ponte_da_planilha_diz_sem_destino_e_nao_balde():
    ws = _abas(_excel(FIXTURES / "sintetica_integra.xls"))[ABA_CONFERENCIA]
    textos = _textos(ws)
    assert "(-) Débitos sem destino" in textos
    assert "(+) Créditos sem destino" in textos
    assert not any(t.startswith("(") and "balde" in t for t in textos)


def test_secao_dos_itens_sem_destino_sem_a_palavra_balde():
    textos = _textos(_abas(_excel(FIXTURES / "sintetica_integra.xls"))[ABA_CONFERENCIA])
    assert any(t.startswith("Itens sem destino — ordenado por valor") for t in textos)
    assert not any(t.startswith("Balde") for t in textos)


def test_secao_vazia_diz_o_fato():
    textos = _textos(_abas(_excel(GOLDEN))[ABA_CONFERENCIA])
    assert "Nenhum item sem destino." in textos
    assert not any("todo pagamento casou" in t for t in textos)


# === texto de historico nunca vira formula nem erro ===========
# O openpyxl infere o tipo de toda str: "=..." vira formula ('f') e "#N/A" (e
# cia.) vira celula de erro ('e'). A copia fiel deixaria de ser fiel, e formula
# externa (HYPERLINK/WEBSERVICE) rodaria no Excel do conferente.

HISTORICOS_QUE_O_OPENPYXL_INTERPRETA = [
    '=HYPERLINK("http://exemplo.invalid","x")',
    "=2+3",
    "#N/A",
]


def _razao_xls_com_historicos(caminho: Path, historicos: list[str], *,
                              lote=None, cta_part=38.0,
                              nome_conta="FORNECEDOR FICTICIO LTDA") -> Path:
    """.xls sintetico no layout do Dominio: 1 compra e um debito de 10,00 por
    historico. `lote` e `cta_part` trocam o numero do lote e a Cta.C.Part. do
    1o debito; `nome_conta` e o nome da conta. Dados ficticios."""
    import xlwt

    from tests.fixtures.gerar_sinteticas import CNPJ, COL_SALDO, EMPRESA, PERIODO, _serial
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Razão")
    ws.write(0, 0, "Empresa:"); ws.write(0, 2, EMPRESA)
    ws.write(1, 0, "C.N.P.J.:"); ws.write(1, 2, CNPJ)
    ws.write(2, 0, "Período:"); ws.write(2, 2, PERIODO)
    ws.write(3, 0, "RAZÃO")
    ws.write(5, 0, "Data"); ws.write(5, 1, "Número"); ws.write(5, 2, "Histórico")
    ws.write(5, 7, "Cta.C.Part."); ws.write(5, 8, "Débito"); ws.write(5, 9, "Crédito")
    ws.write(5, COL_SALDO, "Saldo-Exercício")
    ws.write(6, 0, "Conta:"); ws.write(6, 1, 7001.0)
    ws.write(6, 2, "2.01.01.001.00"); ws.write(6, 5, nome_conta)
    ws.write(7, 2, "SALDO ANTERIOR"); ws.write(7, COL_SALDO, 0.0)
    linhas = [("Vr. compras conf. 1234 FORNECEDOR FICTICIO LTDA", None, 1000.0)]
    linhas += [(h, 10.0, None) for h in historicos]
    r, saldo, td, tc = 8, 0.0, 0.0, 0.0
    for i, (hist, deb, cred) in enumerate(linhas):
        # o 1o debito leva o lote e a Cta.C.Part. pedidos
        lote_da_linha = lote if (i == 1 and lote is not None) else float(5000 + i)
        cta_da_linha = cta_part if i == 1 else 38.0
        ws.write(r, 0, _serial(date(2025, 3, 1 + i))); ws.write(r, 1, lote_da_linha)
        ws.write(r, 2, hist); ws.write(r, 7, cta_da_linha)
        if deb is not None:
            ws.write(r, 8, deb); saldo += deb; td += deb
        if cred is not None:
            ws.write(r, 9, cred); saldo -= cred; tc += cred
        ws.write(r, COL_SALDO, saldo); r += 1
    ws.write(r + 1, 4, "Total da conta:"); ws.write(r + 1, 8, td); ws.write(r + 1, 9, tc)
    wb.save(str(caminho))
    return caminho


def test_historico_que_parece_formula_ou_erro_sai_como_texto_nas_duas_abas(tmp_path):
    historicos = HISTORICOS_QUE_O_OPENPYXL_INTERPRETA
    caminho = _razao_xls_com_historicos(tmp_path / "formulas.xls", historicos)
    wb = _abas(_excel(caminho))
    for aba in (ABA_RAZAO, ABA_CONFERENCIA):
        achadas = {c.value: c for c in _celulas(wb[aba])
                   if isinstance(c.value, str) and c.value in historicos}
        assert set(achadas) == set(historicos), aba          # o texto esta la, igual
        for texto, celula in achadas.items():
            assert celula.data_type == "s", (aba, texto, celula.data_type)


def test_nenhuma_celula_da_planilha_vira_formula_ou_erro_no_golden():
    """CONTROLE: o golden nao tem texto de formula; o helper nao muda o tipo
    de nenhuma celula (numeros e datas seguem como antes)."""
    wb = _abas(_excel(GOLDEN))
    tipos = {c.data_type for ws in wb for c in _celulas(ws)}
    assert tipos <= {"s", "n", "d"}


# === caractere de controle no texto nao derruba a exportacao
# O parser aceita o texto; o openpyxl RECUSA na escrita (IllegalCharacterError) e
# a exportacao do lote roda antes do render: a tela inteira caia, com o arquivo
# integro do mesmo lote. Todo texto escrito passa por _grava, que troca o
# controle por espaco.

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # noqa: E402

_PONTOS_COM_CONTROLE = {
    "historico": dict(historicos=["TARIFA\x0cLINHA COLADA"]),
    "numero do lote": dict(historicos=["TARIFA BANCARIA"], lote="50\x01"),
    "Cta.C.Part.": dict(historicos=["TARIFA BANCARIA"], cta_part="38\x07"),
    "nome da conta": dict(historicos=["TARIFA BANCARIA"],
                          nome_conta="FORNECEDOR\x08 FICTICIO"),
    "nome do arquivo": dict(historicos=["TARIFA BANCARIA"],
                            nome_arquivo="arquivo\x1fcontrole.xls"),
}


@pytest.mark.parametrize("ponto", list(_PONTOS_COM_CONTROLE))
def test_caractere_de_controle_no_texto_nao_derruba_a_exportacao(tmp_path, ponto):
    pedido = dict(_PONTOS_COM_CONTROLE[ponto])
    nome_arquivo = pedido.pop("nome_arquivo", "controle.xls")
    historicos = pedido.pop("historicos")
    caminho = _razao_xls_com_historicos(tmp_path / "controle.xls", historicos, **pedido)
    dados = _excel(caminho, nome_arquivo)
    wb = _abas(dados)
    for ws in wb:
        ruins = [c.coordinate for c in _celulas(ws)
                 if isinstance(c.value, str) and ILLEGAL_CHARACTERS_RE.search(c.value)]
        assert ruins == [], (ponto, ws.title)


def test_controle_no_historico_vira_espaco_nas_duas_abas(tmp_path):
    caminho = _razao_xls_com_historicos(tmp_path / "c.xls", ["TARIFA\x0cLINHA COLADA"])
    wb = _abas(_excel(caminho))
    for aba in (ABA_RAZAO, ABA_CONFERENCIA):
        assert "TARIFA LINHA COLADA" in _textos(wb[aba]), aba


def test_controle_no_nome_do_arquivo_vira_espaco_no_cabecalho(tmp_path):
    caminho = _razao_xls_com_historicos(tmp_path / "c.xls", ["TARIFA BANCARIA"])
    textos = _textos(_abas(_excel(caminho, "arquivo\x1fcontrole.xls"))[ABA_CONFERENCIA])
    assert any("arquivo controle.xls" in t for t in textos)


def test_texto_sem_controle_sai_identico():
    """CONTROLE: tab, quebra de linha e acento passam (o openpyxl aceita) e nada muda."""
    ws = openpyxl.Workbook().active
    from relatorio.excel import _grava
    for texto in ["Ação\tcom tab", "linha1\nlinha2", "R$ 1.000,00 — fim"]:
        assert _grava(ws, 1, 1, texto).value == texto


def test_formula_com_controle_no_meio_sai_como_texto_e_sem_o_controle(tmp_path):
    """Os dois juntos: '=2<controle>+3' nao derruba a exportacao, o
    controle vira espaco e a celula segue TEXTO, nao formula."""
    caminho = _razao_xls_com_historicos(tmp_path / "fc.xls", ["=2\x0b+3"])
    wb = _abas(_excel(caminho))
    for aba in (ABA_RAZAO, ABA_CONFERENCIA):
        achadas = [c for c in _celulas(wb[aba]) if c.value == "=2 +3"]
        assert achadas, aba
        assert all(c.data_type == "s" for c in achadas), aba


# === planilha com U+FFFF/U+FFFE (e substituto solto) no texto tem de REABRIR ==========
# O XML do xlsx nao aceita U+D800-U+DFFF (2.048), U+FFFE e U+FFFF (so esses 2.050 nao reabrem; medido por varredura
# de todos os codigos). O _grava ja troca o controle por espaco; passa a trocar tambem esses.

_NAO_REABRE = ["\uffff", "\ufffe"]
_CLASSE_QUE_NAO_REABRE = ["\ud800", "\udfff", "\ufffe", "\uffff"]
_PONTOS_COM_CLASSE = {
    ponto: pedido for ponto, pedido in {
        "historico": dict(historicos=["TARIFA{c}LINHA COLADA"]),
        "numero do lote": dict(historicos=["TARIFA BANCARIA"], lote="50{c}"),
        "Cta.C.Part.": dict(historicos=["TARIFA BANCARIA"], cta_part="38{c}"),
        "nome da conta": dict(historicos=["TARIFA BANCARIA"], nome_conta="FORNECEDOR{c} FICTICIO"),
        "nome do arquivo": dict(historicos=["TARIFA BANCARIA"], nome_arquivo="arquivo{c}classe.xls"),
    }.items()}


@pytest.mark.parametrize("c", _NAO_REABRE)
def test_u_ffff_e_u_fffe_no_historico_a_planilha_reabre_com_espaco(tmp_path, c):
    caminho = _razao_xls_com_historicos(tmp_path / "n.xls", [f"TARIFA{c}LINHA COLADA"])
    wb = _abas(_excel(caminho))                                   # RED: ParseError ao reabrir
    for aba in (ABA_RAZAO, ABA_CONFERENCIA):
        assert "TARIFA LINHA COLADA" in _textos(wb[aba]), aba


@pytest.mark.parametrize("c", _NAO_REABRE)
@pytest.mark.parametrize("ponto", list(_PONTOS_COM_CLASSE))
def test_os_5_pontos_de_texto_reabrem_e_nao_guardam_o_caractere(tmp_path, ponto, c):
    pedido = {k: (v.replace("{c}", c) if isinstance(v, str) else [x.replace("{c}", c) for x in v])
              for k, v in _PONTOS_COM_CLASSE[ponto].items()}
    nome_arquivo = pedido.pop("nome_arquivo", "n.xls")
    historicos = pedido.pop("historicos")
    caminho = _razao_xls_com_historicos(tmp_path / "n.xls", historicos, **pedido)
    wb = _abas(_excel(caminho, nome_arquivo))
    for ws in wb:
        ruins = [x.coordinate for x in _celulas(ws)
                 if isinstance(x.value, str) and any(ch in x.value for ch in ("\ufffe", "\uffff"))]
        assert ruins == [], (ponto, ws.title)


@pytest.mark.parametrize("c", _CLASSE_QUE_NAO_REABRE)
def test_grava_substitui_a_classe_por_um_espaco(c):
    from relatorio.excel import _grava
    ws = openpyxl.Workbook().active
    assert _grava(ws, 1, 1, "A" + c + "B").value == "A B"


@pytest.mark.parametrize("texto", ["\ufffd", "\ufdd0", "\U0001fffe", "\U0010ffff", "", "\ud7ff", "\ue000",
                                   "Ação\tcom tab", "linha1\nlinha2", "\U0001f600"])
def test_controle_o_que_reabre_hoje_sai_identico(texto):
    from relatorio.excel import _grava
    ws = openpyxl.Workbook().active
    assert _grava(ws, 1, 1, texto).value == texto


@pytest.mark.parametrize("c", _NAO_REABRE)
def test_a_planilha_e_igual_celula_a_celula_a_da_mesma_razao_com_espaco(tmp_path, c):
    com = _razao_xls_com_historicos(tmp_path / "com.xls", [f"TARIFA {c} SEM PADRAO"])
    sem = _razao_xls_com_historicos(tmp_path / "sem.xls", ["TARIFA   SEM PADRAO"])
    wb_com, wb_sem = _abas(_excel(com, "mesmo.xls")), _abas(_excel(sem, "mesmo.xls"))
    for aba in (ABA_RAZAO, ABA_CONFERENCIA):
        a = [(x.coordinate, x.value) for x in _celulas(wb_com[aba])]
        b = [(x.coordinate, x.value) for x in _celulas(wb_sem[aba])]
        assert a == b, aba
